"""Excel connector — D18.

Reads mapped cell values from an .xlsx workbook. Pure read:
  - openpyxl with read_only=True, data_only=True, keep_vba=False
  - no macro execution, no formula evaluation (cached values only), no external-link following
  - the uploaded bytes are never written back

Returns {(worksheet, cell): value} for the requested mapping. Raises typed IngestionErrors
for unreadable workbooks, missing worksheets/cells, and formulas without a cached value.
"""

from __future__ import annotations

import hashlib
import io
import zipfile
from typing import Any

import openpyxl
from openpyxl.utils import range_boundaries
from openpyxl.utils.exceptions import InvalidFileException

from backend.app.ingestion.errors import (
    CellNotFoundError,
    FormulaWithoutCachedValueError,
    WorkbookUnreadableError,
    WorksheetNotFoundError,
)
from backend.app.ingestion.mappings import SourceMapping

XLSX_SUFFIX = ".xlsx"


def sha256_hex(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _looks_like_formula(value: Any) -> bool:
    return isinstance(value, str) and value.startswith("=")


def read_cells(workbook_bytes: bytes, mapping: SourceMapping) -> dict[tuple[str, str], Any]:
    """
    Return {(worksheet, cell): value} for every cell the mapping references.

    Two loads: data_only=True yields each formula's cached value (we never evaluate formulas),
    data_only=False lets us tell a genuinely blank cell from a formula whose value was never
    cached. A formula with no cached value is rejected with FormulaWithoutCachedValueError.
    """
    values_wb = _load(workbook_bytes, data_only=True)
    try:
        formulas_wb = _load(workbook_bytes, data_only=False)
        try:
            out: dict[tuple[str, str], Any] = {}
            sheet_names = set(values_wb.sheetnames)
            for cm in mapping.cells:
                if cm.worksheet not in sheet_names:
                    raise WorksheetNotFoundError(
                        f"Worksheet {cm.worksheet!r} not found; available: {sorted(sheet_names)}"
                    )
                value = _read_one(values_wb[cm.worksheet], cm.worksheet, cm.cell)
                if value is None:
                    raw = _read_one(formulas_wb[cm.worksheet], cm.worksheet, cm.cell)
                    if _looks_like_formula(raw):
                        raise FormulaWithoutCachedValueError(
                            f"Cell {cm.worksheet}!{cm.cell} holds a formula with no cached value; "
                            "re-save the workbook in Excel so values are cached "
                            "(formulas are not evaluated)."
                        )
                out[(cm.worksheet, cm.cell)] = value
            return out
        finally:
            formulas_wb.close()
    finally:
        values_wb.close()


def _load(workbook_bytes: bytes, *, data_only: bool):
    try:
        return openpyxl.load_workbook(
            io.BytesIO(workbook_bytes), read_only=True, data_only=data_only, keep_vba=False,
        )
    except (InvalidFileException, zipfile.BadZipFile, KeyError, OSError) as exc:
        raise WorkbookUnreadableError(f"Workbook could not be read: {exc}") from exc


def _read_one(ws, sheet: str, cell: str) -> Any:
    try:
        min_col, min_row, max_col, max_row = range_boundaries(cell)
    except (ValueError, TypeError) as exc:
        raise CellNotFoundError(f"Invalid cell reference {sheet}!{cell}: {exc}") from exc
    # A single-cell reference has coincident boundaries.
    if (min_col, min_row) != (max_col, max_row):
        raise CellNotFoundError(f"Mapping cell {sheet}!{cell} must be a single cell, not a range.")
    try:
        return ws.cell(row=min_row, column=min_col).value
    except (IndexError, ValueError) as exc:
        raise CellNotFoundError(f"Cell {sheet}!{cell} is out of range: {exc}") from exc
