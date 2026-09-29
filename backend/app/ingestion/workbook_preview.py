"""Guided Excel workflow support — D24.

D18 gave the platform a working Excel connector, but the only thing you could do with a
workbook was *commit* it: upload, and the value was written, change-detected, and
propagated in one irreversible step. There was no way to ask "what is in this file, what
would it change, and is it even valid?" before committing, and no way to obtain a
correctly-shaped workbook in the first place.

This module adds the two missing halves of the loop, both **read-only**:

  * :func:`build_template`  — generate a real .xlsx from a named mapping, pre-filled with
    the dataset's current values, so a non-developer starts from a correct workbook rather
    than guessing cell addresses.
  * :func:`preview_workbook` — dry-run an uploaded workbook: read every mapped cell,
    compare it to the current value, and run the real contract validation — **without
    writing anything, without creating an IngestionRun, and without propagating**.

Both reuse the D18 connector and the D16 contract layer unchanged. `preview_workbook`
calls exactly the same `validate_record` the commit path calls, so a preview that says
"valid" and a commit that succeeds cannot disagree about the rules.

Nothing here evaluates formulas, runs VBA, or executes macros — the D18 connector reads
cached values only, and that limitation is reported to the caller rather than hidden.
"""

from __future__ import annotations

import io
from dataclasses import asdict, dataclass
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy.orm import Session

from backend.app.ingestion.errors import IngestionError, UnknownMappingError
from backend.app.ingestion.excel_connector import read_cells, sha256_hex
from backend.app.ingestion.mappings import SourceMapping, get_mapping
from backend.app.persistence.dataset_repository import DatasetRepository
from backend.app.persistence.exceptions import NotFoundError
from backend.app.services.contract_validation import ContractViolationError
from backend.app.services.data_contract_manager import DataContractManager
from backend.app.services.dataset_values import DatasetValueService
from backend.app.services.federation import FederationService

# A workbook cell whose numeric value differs from the current one by less than this is
# treated as unchanged. It matches ordinary float round-tripping through Excel, not a
# tolerance on the model itself.
_FLOAT_EPSILON = 1e-9


@dataclass
class FieldPreview:
    """One mapped cell, fully explained for a non-developer."""

    worksheet: str
    cell: str                    # e.g. "B2" — where the value came from
    dataset: str                 # target dataset name
    field: str                   # target contract field name
    unit: str | None             # unit from the dataset's active contract
    expected_type: str           # what the mapping declares the cell should hold
    current_value: Any           # what the platform holds now (None when never set)
    new_value: Any               # what this workbook says
    changed: bool
    valid: bool
    message: str | None          # why it is invalid, or a note (e.g. "unchanged")


@dataclass
class WorkbookPreview:
    """The complete dry-run result. Nothing was written to produce this."""

    mapping_key: str
    mapping_description: str
    source_name: str
    content_sha256: str
    fields: list[FieldPreview]
    valid: bool                  # every field valid → commit will not be rejected
    changed_fields: list[str]    # fields whose value would change
    would_change: bool
    datasets: list[str]
    errors: list[str]            # workbook-level failures (bad sheet, formula, etc.)
    notes: list[str]             # honest limitations that apply to every ingestion

    def to_dict(self) -> dict[str, Any]:
        d = asdict(self)
        d["fields"] = [asdict(f) if not isinstance(f, dict) else f for f in self.fields]
        return d


CONNECTOR_NOTES = [
    "Formulas are never evaluated. The connector reads the value Excel cached the last "
    "time the workbook was saved, so a formula cell must have been saved by Excel at "
    "least once or the upload is rejected.",
    "Macros and VBA are never executed. Workbooks are opened read-only with macro "
    "support disabled.",
    "Only the cells listed in the mapping are read. Everything else in the workbook — "
    "other sheets, formatting, charts, extra rows — is ignored.",
    "The uploaded file is never written back to, and never stored as a file; only its "
    "SHA-256 hash and the mapped values are recorded.",
]


# ---------------------------------------------------------------------------
# Field metadata
# ---------------------------------------------------------------------------

def _contract_units(db: Session, mapping: SourceMapping) -> dict[tuple[str, str], str | None]:
    """(dataset name, field name) → unit, taken from each dataset's active contract."""
    datasets = DatasetRepository(db)
    contracts = DataContractManager(db)
    units: dict[tuple[str, str], str | None] = {}
    for dataset_name in mapping.datasets():
        try:
            ds = datasets.get_by_name(dataset_name)
        except NotFoundError:
            continue
        schema = contracts.get_active_schema(ds.id)
        if schema is None:
            continue
        for spec in schema.fields:
            units[(dataset_name, spec.name)] = spec.unit
    return units


def _current_values(db: Session, mapping: SourceMapping) -> dict[str, dict[str, Any]]:
    """dataset name → its current record ({} when the dataset has never been written)."""
    datasets = DatasetRepository(db)
    values = DatasetValueService(db)
    out: dict[str, dict[str, Any]] = {}
    for dataset_name in mapping.datasets():
        try:
            ds = datasets.get_by_name(dataset_name)
        except NotFoundError:
            out[dataset_name] = {}
            continue
        current = values.get_value(ds.id)
        out[dataset_name] = current if isinstance(current, dict) else {}
    return out


def _is_changed(current: Any, new: Any) -> bool:
    if current is None and new is None:
        return False
    if isinstance(current, (int, float)) and isinstance(new, (int, float)) \
            and not isinstance(current, bool) and not isinstance(new, bool):
        return abs(float(current) - float(new)) > _FLOAT_EPSILON
    return current != new


def mapping_preview(db: Session, mapping_key: str) -> dict[str, Any]:
    """Describe a mapping without any workbook: which cell feeds which field, in what
    unit, and what the platform currently holds for it.

    This is what the UI shows *before* a user has uploaded anything, so the cell/field
    correspondence is visible up front rather than discovered by trial and error.
    """
    mapping = get_mapping(mapping_key)
    if mapping is None:
        raise UnknownMappingError(f"Unknown mapping key {mapping_key!r}.")

    units = _contract_units(db, mapping)
    current = _current_values(db, mapping)
    return {
        "key": mapping.key,
        "description": mapping.description,
        "datasets": mapping.datasets(),
        "fields": [
            {
                "worksheet": cm.worksheet,
                "cell": cm.cell,
                "dataset": cm.target_dataset,
                "field": cm.target_field,
                "unit": units.get((cm.target_dataset, cm.target_field)),
                "expected_type": cm.expected_type,
                "current_value": current.get(cm.target_dataset, {}).get(cm.target_field),
            }
            for cm in mapping.cells
        ],
        "notes": list(CONNECTOR_NOTES),
    }


# ---------------------------------------------------------------------------
# Template generation
# ---------------------------------------------------------------------------

def build_template(db: Session, mapping_key: str) -> tuple[bytes, str]:
    """Generate a ready-to-edit .xlsx for a mapping, pre-filled with current values.

    The produced workbook is exactly what the connector expects: values sit in the mapped
    cells, and the surrounding labels/units/help are placed in columns the mapping does not
    read. Returns (bytes, suggested filename).
    """
    mapping = get_mapping(mapping_key)
    if mapping is None:
        raise UnknownMappingError(f"Unknown mapping key {mapping_key!r}.")

    units = _contract_units(db, mapping)
    current = _current_values(db, mapping)

    wb = Workbook()
    # Remove the default sheet; sheets are created per mapped worksheet below.
    wb.remove(wb.active)

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="2563EB")
    note_font = Font(italic=True, size=9, color="475569")

    worksheets = []
    for cm in mapping.cells:
        if cm.worksheet not in worksheets:
            worksheets.append(cm.worksheet)

    for sheet_name in worksheets:
        ws = wb.create_sheet(sheet_name)
        ws["A1"] = "Assumption"
        ws["B1"] = "Value  ← edit this column"
        ws["C1"] = "Unit"
        ws["D1"] = "Field name (do not edit)"
        for col in ("A1", "B1", "C1", "D1"):
            ws[col].font = header_font
            ws[col].fill = header_fill
            ws[col].alignment = Alignment(vertical="center")

        for cm in mapping.cells:
            if cm.worksheet != sheet_name:
                continue
            row = int("".join(ch for ch in cm.cell if ch.isdigit()))
            label = cm.target_field.replace("_", " ").capitalize()
            ws.cell(row=row, column=1, value=label)
            ws[cm.cell] = current.get(cm.target_dataset, {}).get(cm.target_field)
            ws.cell(row=row, column=3, value=units.get((cm.target_dataset, cm.target_field)))
            ws.cell(row=row, column=4, value=cm.target_field)

        last_row = max(
            int("".join(ch for ch in cm.cell if ch.isdigit()))
            for cm in mapping.cells if cm.worksheet == sheet_name
        )
        for offset, note in enumerate(CONNECTOR_NOTES, start=2):
            cell = ws.cell(row=last_row + offset, column=1, value=f"Note: {note}")
            cell.font = note_font
        ws.column_dimensions["A"].width = 38
        ws.column_dimensions["B"].width = 22
        ws.column_dimensions["C"].width = 16
        ws.column_dimensions["D"].width = 34

    # A second sheet documenting the mapping itself. The connector never reads it.
    doc = wb.create_sheet("How to use")
    doc["A1"] = f"Mapping: {mapping.key}"
    doc["A1"].font = Font(bold=True, size=13)
    doc["A2"] = mapping.description
    lines = [
        "",
        "1. Edit only the Value column on the assumptions sheet.",
        "2. Do not insert or delete rows — the platform reads fixed cell addresses.",
        "3. Save as .xlsx (not .xls, .xlsm or .csv).",
        "4. Upload it under Excel Assumptions in the app and press Validate.",
        "5. Validate is a dry run: it shows what would change and never writes anything.",
        "6. Commit ingestion writes the values, records provenance, and propagates the",
        "   change through every dependent model.",
        "",
        "Cell map (the only cells that are read):",
    ]
    for i, line in enumerate(lines, start=4):
        doc.cell(row=i, column=1, value=line)
    start = 4 + len(lines)
    doc.cell(row=start, column=1, value="Sheet")
    doc.cell(row=start, column=2, value="Cell")
    doc.cell(row=start, column=3, value="Dataset")
    doc.cell(row=start, column=4, value="Field")
    doc.cell(row=start, column=5, value="Unit")
    for col in range(1, 6):
        doc.cell(row=start, column=col).font = header_font
        doc.cell(row=start, column=col).fill = header_fill
    for offset, cm in enumerate(mapping.cells, start=1):
        doc.cell(row=start + offset, column=1, value=cm.worksheet)
        doc.cell(row=start + offset, column=2, value=cm.cell)
        doc.cell(row=start + offset, column=3, value=cm.target_dataset)
        doc.cell(row=start + offset, column=4, value=cm.target_field)
        doc.cell(row=start + offset, column=5, value=units.get((cm.target_dataset, cm.target_field)))
    for col, width in zip("ABCDE", (26, 10, 24, 34, 16)):
        doc.column_dimensions[col].width = width

    buffer = io.BytesIO()
    wb.save(buffer)
    wb.close()
    return buffer.getvalue(), f"{mapping.key}_template.xlsx"


# ---------------------------------------------------------------------------
# Dry-run preview
# ---------------------------------------------------------------------------

def preview_workbook(db: Session, *, source_name: str, workbook_bytes: bytes,
                     mapping_key: str) -> WorkbookPreview:
    """Read an uploaded workbook and report what committing it would do. Writes nothing.

    Runs the real contract validation used by the commit path, so "valid: true" here means
    the commit will not be rejected by the contract layer.
    """
    mapping = get_mapping(mapping_key)
    if mapping is None:
        raise UnknownMappingError(f"Unknown mapping key {mapping_key!r}.")

    content_sha256 = sha256_hex(workbook_bytes)
    units = _contract_units(db, mapping)
    current = _current_values(db, mapping)
    errors: list[str] = []

    empty = WorkbookPreview(
        mapping_key=mapping.key, mapping_description=mapping.description,
        source_name=source_name, content_sha256=content_sha256, fields=[], valid=False,
        changed_fields=[], would_change=False, datasets=mapping.datasets(),
        errors=errors, notes=list(CONNECTOR_NOTES),
    )

    if not source_name.lower().endswith(".xlsx"):
        errors.append(f"Only .xlsx files are accepted; got {source_name!r}.")
        return empty

    try:
        cells = read_cells(workbook_bytes, mapping)
    except IngestionError as exc:
        errors.append(str(exc))
        return empty

    # Per-field view.
    fields: list[FieldPreview] = []
    records: dict[str, dict[str, Any]] = {}
    for cm in mapping.cells:
        new_value = cells[(cm.worksheet, cm.cell)]
        cur = current.get(cm.target_dataset, {}).get(cm.target_field)
        records.setdefault(cm.target_dataset, {})[cm.target_field] = new_value
        changed = _is_changed(cur, new_value)
        fields.append(FieldPreview(
            worksheet=cm.worksheet, cell=cm.cell, dataset=cm.target_dataset,
            field=cm.target_field, unit=units.get((cm.target_dataset, cm.target_field)),
            expected_type=cm.expected_type, current_value=cur, new_value=new_value,
            changed=changed, valid=True,
            message=None if changed else "unchanged",
        ))

    by_field = {(f.dataset, f.field): f for f in fields}

    # Ownership + contract validation, reusing the commit path's own rules.
    datasets = DatasetRepository(db)
    contracts = DataContractManager(db)
    federation = FederationService(db)
    for dataset_name, record in records.items():
        try:
            ds = datasets.get_by_name(dataset_name)
        except NotFoundError:
            errors.append(f"Mapping targets unknown dataset {dataset_name!r}.")
            for f in fields:
                if f.dataset == dataset_name:
                    f.valid = False
                    f.message = "target dataset does not exist"
            continue
        if federation.producers_of(ds.id):
            errors.append(
                f"Dataset {dataset_name!r} is produced by a model and cannot be ingested into."
            )
            for f in fields:
                if f.dataset == dataset_name:
                    f.valid = False
                    f.message = "dataset is model-produced; ingestion is not allowed"
            continue
        try:
            contracts.validate_record(ds.id, record, phase="write")
        except ContractViolationError as exc:
            for violation in exc.violations:
                target = by_field.get((dataset_name, violation.field))
                if target is not None:
                    target.valid = False
                    target.message = violation.message
                else:
                    errors.append(f"{dataset_name}.{violation.field}: {violation.message}")
            if not exc.violations:
                errors.append(str(exc))

    valid = not errors and all(f.valid for f in fields)
    changed_fields = [f.field for f in fields if f.changed]
    return WorkbookPreview(
        mapping_key=mapping.key, mapping_description=mapping.description,
        source_name=source_name, content_sha256=content_sha256, fields=fields,
        valid=valid, changed_fields=changed_fields, would_change=bool(changed_fields),
        datasets=mapping.datasets(), errors=errors, notes=list(CONNECTOR_NOTES),
    )
