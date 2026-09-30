"""D25 Stage 1 — React product foundation.

Covers the product read models the React app is built on (workspace, catalog, federation
map, semantic activity, scenario explanation, run comparison, upload impact), the React
bundle serving at /app, readiness, the rollback switch to the legacy UI, workbook archive
hardening, and the shared-database configuration helpers.

Every expected figure is produced by the real D17 domain through the real engine; nothing
here is asserted against a hard-coded UI string except sentence-case labels, which ARE the
requirement.
"""

from __future__ import annotations

import io
import json
import re
import zipfile

import openpyxl
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine, event as sa_event
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from backend.app import main as main_module
from backend.app.api.deps import get_db
from backend.app.config import settings
from backend.app.config.settings import _db_host_port, _normalize_db_url
from backend.app.ingestion.errors import WorkbookUnreadableError
from backend.app.ingestion.excel_connector import check_archive
from backend.app.main import api
from backend.app.persistence.database import Base
from backend.app.product.labels import display_unit, humanize
from backend.app.services.contract_validation import parse_schema
from backend.app.ui import decision_state
from backend.app.ui.port_decision_seed import seed_port_decision
from tests.demo_support import db_override, seed_demo

UUID = re.compile(r"^[0-9a-f]{8}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{12}$")


def _engine():
    eng = create_engine("sqlite:///:memory:", connect_args={"check_same_thread": False},
                        poolclass=StaticPool)

    @sa_event.listens_for(eng, "connect")
    def _fk(conn, _):
        conn.execute("PRAGMA foreign_keys=ON")

    Base.metadata.create_all(bind=eng)
    return eng


@pytest.fixture
def product():
    eng = _engine()
    SF = sessionmaker(bind=eng)
    seed_demo(SF)                       # the D0–D16 fuel-price chain coexists, as in the app
    db = SF()
    cfg = seed_port_decision(db)
    db.commit()
    db.close()
    api.dependency_overrides[get_db] = db_override(SF)
    with TestClient(api) as client:
        orig = decision_state.config
        decision_state.config = cfg
        try:
            yield client, cfg
        finally:
            decision_state.config = orig
    api.dependency_overrides.clear()
    eng.dispose()


def _workbook(client, bunker: float) -> bytes:
    raw = client.get("/api/ingestions/template?mapping=port_assumptions_v1").content
    wb = openpyxl.load_workbook(io.BytesIO(raw))
    wb["Assumptions"]["B2"] = bunker
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Labels and sentence case
# ---------------------------------------------------------------------------

class TestLabels:
    @pytest.mark.parametrize("ident, label", [
        ("annual_fuel_cost_usd", "Annual fuel cost"),
        ("fuel_cost_per_teu", "Fuel cost per TEU"),
        ("annual_emissions_tco2", "Annual emissions"),
        ("utilization_pct", "Utilization"),
        ("bunker_price", "Bunker price"),
        ("annualTeuForecast", "Annual TEU forecast"),
    ])
    def test_humanize_is_sentence_case_and_drops_unit_suffixes(self, ident, label):
        assert humanize(ident) == label

    def test_display_unit_uses_subscript(self):
        assert display_unit("tCO2/year") == "tCO₂/year"

    def test_contract_label_is_presentation_only(self):
        schema = parse_schema({"fields": [
            {"name": "x", "type": "float", "label": "Bunker price", "unit": "USD/t"},
        ]})
        assert schema.fields[0].label == "Bunker price"
        # Unlabelled contracts still parse exactly as before.
        assert parse_schema({"fields": [{"name": "y", "type": "float"}]}).fields[0].label is None

    def test_seeded_names_are_sentence_case(self, product):
        client, _ = product
        ws = client.get("/api/workspace").json()
        assert [b["name"] for b in ws["baselines"]] == ["Current operations"]
        assert [s["name"] for s in ws["scenarios"]] == ["Higher bunker price"]
        names = {m["name"] for m in client.get("/api/federation/map").json()["models"]}
        for name in names:
            words = name.split()
            assert all(not w[0].isupper() for w in words[1:] if w not in ("TEU",)), name


# ---------------------------------------------------------------------------
# Workspace, explanation, comparison
# ---------------------------------------------------------------------------

class TestWorkspace:
    def test_default_selection_is_derived_from_records(self, product):
        client, cfg = product
        ws = client.get("/api/workspace").json()
        assert ws["default"] == {"baseline_id": cfg["baseline_id"], "scenario_id": cfg["scenario_id"]}
        (ov,) = ws["scenarios"][0]["overrides"]
        assert ov["field_label"] == "Bunker price"
        assert ov["dataset_label"] == "Port assumptions"
        assert ov["unit"] == "USD/t"
        # The baseline value is what the baseline run consumed, read from its step snapshot.
        assert ov["baseline_value_known"] is True
        assert ov["baseline_value"] == 600.0
        assert ov["value"] == 750.0

    def test_assumptions_carry_labels_units_and_provenance(self, product):
        client, _ = product
        ws = client.get("/api/workspace").json()
        port = next(d for d in ws["assumptions"] if d["name"] == "port_assumptions")
        fields = {f["name"]: f for f in port["fields"]}
        assert fields["bunker_price"]["label"] == "Bunker price"
        assert fields["number_of_berths"]["min"] == 1.0
        assert port["current_source"]["source_type"] == "seed"


class TestExplanation:
    def test_field_level_causal_path(self, product):
        client, cfg = product
        ex = client.get(f"/api/scenarios/{cfg['scenario_id']}/explanation").json()
        assert [c["field_label"] for c in ex["changes"]] == ["Bunker price"]
        assert ex["changes"][0]["baseline_value"] == 600.0
        assert ex["changes"][0]["scenario_value"] == 750.0
        assert [p["model_name"] for p in ex["path"]] == ["Port fuel cost", "Decision summary"]
        fuel = ex["path"][0]
        assert [r["field_label"] for r in fuel["reads_changed"]] == ["Bunker price"]
        assert {o["field_label"] for o in fuel["outputs_changed"]} == {"Annual fuel cost", "Fuel cost per TEU"}
        assert [o["field_label"] for o in fuel["outputs_unchanged"]] == ["Annual fuel burned"]
        assert {u["model_name"] for u in ex["unaffected"]} == {"Throughput", "Berth utilization", "Port emissions"}

    def test_comparison_values_come_from_results(self, product):
        client, cfg = product
        ex = client.get(f"/api/scenarios/{cfg['scenario_id']}/explanation").json()
        m = next(x for x in ex["comparison"]["metrics"]
                 if x["terminal"] and x["field"] == "annual_fuel_cost_usd")
        assert (m["baseline"], m["scenario"]) == (12_000_000.0, 15_000_000.0)
        assert m["field_label"] == "Annual fuel cost" and m["unit"] == "USD/year"
        assert m["dataset_label"] == "Decision summary"

    def test_unexecuted_scenario_is_409_not_500(self, product):
        client, cfg = product
        new = client.post("/api/scenarios", json={"baseline_id": cfg["baseline_id"], "name": "Unrun"}).json()
        r = client.get(f"/api/scenarios/{new['id']}/explanation")
        assert r.status_code == 409

    def test_run_compare_generic_and_404(self, product):
        client, cfg = product
        r = client.get(f"/api/runs/compare?base={cfg['baseline_run_id']}&target={cfg['scenario_run_id']}")
        assert r.status_code == 200 and r.json()["changed_count"] == 4
        assert client.get(f"/api/runs/compare?base=nope&target={cfg['scenario_run_id']}").status_code == 404

    def test_override_can_be_removed(self, product):
        client, cfg = product
        ov = client.get(f"/api/scenarios/{cfg['scenario_id']}/overrides").json()[0]
        assert client.delete(f"/api/scenarios/{cfg['scenario_id']}/overrides/{ov['id']}").status_code == 204
        assert client.get(f"/api/scenarios/{cfg['scenario_id']}/overrides").json() == []
        assert client.delete(f"/api/scenarios/{cfg['scenario_id']}/overrides/{ov['id']}").status_code == 422


# ---------------------------------------------------------------------------
# Semantic activity
# ---------------------------------------------------------------------------

class TestActivity:
    def test_scenario_and_baseline_rows_are_semantic(self, product):
        client, _ = product
        items = client.get("/api/activity").json()
        sc = next(i for i in items if i["kind"] == "scenario_run")
        assert (sc["title"], sc["subject"], sc["status"]) == ("Scenario run", "Higher bunker price", "succeeded")
        assert sc["headline"]["label"] == "Annual fuel cost"
        assert (sc["headline"]["from"], sc["headline"]["to"]) == (12_000_000.0, 15_000_000.0)
        bl = next(i for i in items if i["kind"] == "baseline_run")
        assert (bl["title"], bl["subject"]) == ("Baseline run", "Current operations")
        assert bl["headline"]["label"] == "Annual throughput"
        assert bl["headline"]["value"] == 20_000_000.0

    def test_no_row_leads_with_an_identifier(self, product):
        client, _ = product
        for item in client.get("/api/activity").json():
            assert not UUID.match(item["subject"]), item
            assert not UUID.match(item["title"]), item
            assert item["technical"], "identifiers must still be available"

    def test_upload_becomes_one_dataset_change_row_with_its_effect(self, product):
        client, _ = product
        files = {"file": ("port.xlsx", _workbook(client, 825.0),
                          "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        out = client.post("/api/ingestions/excel?mapping=port_assumptions_v1", files=files).json()
        assert out["status"] == "ingested"
        items = client.get("/api/activity").json()
        ch = next(i for i in items if i["kind"] == "dataset_change" and i["title"] == "Dataset change")
        assert ch["subject"] == "Bunker price"
        assert ch["context"] == "Port assumptions"
        assert (ch["headline"]["from"], ch["headline"]["to"], ch["headline"]["unit"]) == (600.0, 825.0, "USD/t")
        assert ch["via"]["file_name"] == "port.xlsx"
        assert ch["effect"]["label"] == "Annual fuel cost"
        assert ch["effect"]["to"] == 16_500_000.0
        # The propagation run is folded into the change, not repeated as its own row.
        assert not any(i["kind"] == "model_run" and i["links"]["run_id"] == out["graph_run_id"] for i in items)

    def test_rejected_upload_is_its_own_row(self, product):
        client, _ = product
        files = {"file": ("bad.xlsx", _workbook(client, -5.0),
                          "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        assert client.post("/api/ingestions/excel?mapping=port_assumptions_v1", files=files).json()["status"] == "rejected"
        up = next(i for i in client.get("/api/activity").json() if i["kind"] == "upload")
        assert (up["title"], up["subject"], up["status"]) == ("Workbook upload", "bad.xlsx", "rejected")
        assert up["error"]


class TestUploadImpact:
    def test_impact_lists_changes_models_and_comparison(self, product):
        client, _ = product
        files = {"file": ("port.xlsx", _workbook(client, 900.0),
                          "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        out = client.post("/api/ingestions/excel?mapping=port_assumptions_v1", files=files).json()
        imp = client.get(f"/api/ingestions/{out['ingestion_id']}/impact").json()
        assert [(c["field_label"], c["from"], c["to"]) for c in imp["changes"]] == [("Bunker price", 600.0, 900.0)]
        assert imp["run"]["status"] == "succeeded"
        assert "Port fuel cost" in [m["model_name"] for m in imp["models_run"]]
        moved = {(m["dataset_label"], m["field"]) for m in imp["comparison"]["metrics"] if m["changed"]}
        assert ("Decision summary", "annual_fuel_cost_usd") in moved
        assert imp["ingestion"]["content_sha256"]
        assert client.get("/api/ingestions/nope/impact").status_code == 404


# ---------------------------------------------------------------------------
# Serving the React bundle, readiness, rollback
# ---------------------------------------------------------------------------

@pytest.fixture
def fake_dist(tmp_path, monkeypatch):
    (tmp_path / "assets").mkdir()
    (tmp_path / "index.html").write_text("<!doctype html><div id=root></div>", encoding="utf-8")
    (tmp_path / "assets" / "index-abc.js").write_text("console.log(1)", encoding="utf-8")
    (tmp_path / "build-meta.json").write_text(json.dumps({"build_id": "b1"}), encoding="utf-8")
    (tmp_path.parent / "secret.txt").write_text("outside", encoding="utf-8")
    monkeypatch.setattr(settings, "FRONTEND_DIST", str(tmp_path))
    return tmp_path


class TestFrontendServing:
    def test_spa_routes_fall_back_to_index(self, fake_dist):
        with TestClient(api) as c:
            for path in ("/app/", "/app/scenarios", "/app/network/hub"):
                r = c.get(path)
                assert r.status_code == 200 and 'id=root' in r.text
                assert r.headers["cache-control"] == "no-cache"

    def test_hashed_assets_are_immutable_and_missing_assets_404(self, fake_dist):
        with TestClient(api) as c:
            r = c.get("/app/assets/index-abc.js")
            assert r.status_code == 200 and "immutable" in r.headers["cache-control"]
            assert c.get("/app/assets/missing.js").status_code == 404

    def test_path_traversal_cannot_escape_the_bundle(self, fake_dist):
        with TestClient(api) as c:
            for path in ("/app/../secret.txt", "/app/..%2Fsecret.txt", "/app/assets/..%2F..%2Fsecret.txt"):
                r = c.get(path)
                assert "outside" not in r.text

    def test_missing_bundle_is_a_clear_503(self, tmp_path, monkeypatch):
        monkeypatch.setattr(settings, "FRONTEND_DIST", str(tmp_path / "nope"))
        with TestClient(api) as c:
            r = c.get("/app/")
            assert r.status_code == 503 and "has not been built" in r.text
            assert c.get("/api/build-info").json()["frontend"]["built"] is False

    def test_build_info_reports_the_bundle(self, fake_dist):
        with TestClient(api) as c:
            fe = c.get("/api/build-info").json()["frontend"]
            assert fe["built"] is True and fe["build_id"] == "b1"

    def test_root_rollback_switch(self, fake_dist, monkeypatch):
        with TestClient(api) as c:
            assert c.get("/", follow_redirects=False).headers["location"] == "/app/"
            monkeypatch.setattr(main_module, "PRIMARY_UI", "legacy")
            assert c.get("/", follow_redirects=False).headers["location"] == "/ui/start"
            # The legacy pages stay served either way.
            assert c.get("/ui").status_code == 200

    def test_root_falls_back_to_legacy_when_bundle_missing(self, tmp_path, monkeypatch):
        monkeypatch.setattr(settings, "FRONTEND_DIST", str(tmp_path / "nope"))
        with TestClient(api) as c:
            assert c.get("/", follow_redirects=False).headers["location"] == "/ui/start"

    def test_head_requests_are_answered(self):
        with TestClient(api) as c:
            assert c.head("/health").status_code == 200
            assert c.head("/ready").status_code in (200, 503)


class TestReadiness:
    def test_ready_on_create_all_sqlite(self):
        with TestClient(api) as c:
            body = c.get("/ready").json()
            assert body["database"] == "ok"
            assert body["schema"] == "managed_by_create_all"
            assert body["ready"] is True

    def test_migrated_database_must_be_at_head(self, tmp_path, monkeypatch):
        from alembic import command
        from alembic.config import Config

        from backend.app.web import alembic_head, readiness

        url = f"sqlite:///{tmp_path / 'm.db'}"
        cfg = Config("alembic.ini")
        cfg.attributes["sqlalchemy_url"] = url
        command.upgrade(cfg, "005")
        eng = create_engine(url)
        monkeypatch.setattr(settings, "AUTO_CREATE_SCHEMA", False)
        ok, report = readiness(eng)
        assert not ok and report["schema"] == "migrations_pending"
        command.upgrade(cfg, "head")
        ok, report = readiness(eng)
        assert ok and report["schema_revision"] == alembic_head()
        eng.dispose()

    def test_unreachable_database_never_leaks_its_url(self, monkeypatch):
        from backend.app.web import readiness

        eng = create_engine("postgresql+psycopg://user:hunter2@127.0.0.1:1/x",
                            connect_args={"connect_timeout": 1})
        ok, report = readiness(eng)
        assert not ok and report["database"] == "unreachable"
        assert "hunter2" not in json.dumps(report)


# ---------------------------------------------------------------------------
# Workbook hardening
# ---------------------------------------------------------------------------

def _zip(entries: dict[str, bytes]) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in entries.items():
            z.writestr(name, data)
    return buf.getvalue()


class TestArchiveHardening:
    def test_macro_project_is_rejected(self):
        with pytest.raises(WorkbookUnreadableError, match="macro"):
            check_archive(_zip({"[Content_Types].xml": b"x", "xl/vbaProject.bin": b"\0" * 10}))

    def test_decompression_bomb_is_rejected_before_parsing(self):
        bomb = _zip({"xl/worksheets/sheet1.xml": b"0" * (65 * 1024 * 1024)})
        assert len(bomb) < 1024 * 1024       # tiny on the wire…
        with pytest.raises(WorkbookUnreadableError, match="expands"):
            check_archive(bomb)             # …refused before openpyxl sees it

    def test_not_a_zip_is_rejected(self):
        with pytest.raises(WorkbookUnreadableError):
            check_archive(b"not a zip at all")

    def test_macro_upload_is_recorded_as_rejected(self, product):
        client, _ = product
        good = _workbook(client, 600.0)
        src = zipfile.ZipFile(io.BytesIO(good))
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w") as z:
            for item in src.infolist():
                z.writestr(item, src.read(item.filename))
            z.writestr("xl/vbaProject.bin", b"\0" * 64)
        files = {"file": ("macro.xlsx", buf.getvalue(),
                          "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        out = client.post("/api/ingestions/excel?mapping=port_assumptions_v1", files=files).json()
        assert out["status"] == "rejected" and "macro" in out["error"].lower()


# ---------------------------------------------------------------------------
# Shared-database configuration (no connection is made)
# ---------------------------------------------------------------------------

class TestDatabaseConfig:
    def test_supabase_pooler_url_is_normalised(self):
        url = _normalize_db_url(
            "postgresql://postgres.ref:pw@aws-0-x.pooler.supabase.com:6543/postgres?pgbouncer=true")
        assert url.startswith("postgresql+psycopg://")
        assert "pgbouncer" not in url
        assert url.endswith("?sslmode=require")
        assert _db_host_port(url) == ("aws-0-x.pooler.supabase.com", 6543)

    def test_explicit_sslmode_is_respected(self):
        url = _normalize_db_url("postgresql://u:p@db.ref.supabase.co:5432/postgres?sslmode=verify-full")
        assert url.count("sslmode=") == 1 and "verify-full" in url

    def test_non_supabase_and_sqlite_are_untouched(self):
        assert _normalize_db_url("sqlite:///x.db") == "sqlite:///x.db"
        assert "sslmode" not in _normalize_db_url("postgres://u:p@localhost:5432/db")

    def test_password_with_at_sign_does_not_confuse_host_parsing(self):
        assert _db_host_port("postgresql://u:p%40ss@host.example:5433/db") == ("host.example", 5433)

    def test_build_info_never_exposes_the_database_url(self):
        with TestClient(api) as c:
            info = c.get("/api/build-info").json()
            blob = json.dumps(info)
            assert info["database"] in ("sqlite", "postgresql")
            assert "://" not in blob.replace("http://", "").replace("https://", "")
