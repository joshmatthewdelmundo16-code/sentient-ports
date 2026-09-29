"""D18 — Excel Integration tests.

Unit (connector): read, deterministic SHA-256, missing worksheet/cell, blank/wrong-type/
numeric conversion, corrupt workbook, unsupported extension, oversized upload, and the
cached-value / formula policy.

Integration: workbook → mapping → dataset through the D16 contract/change-detection path,
hash + provenance, first ingestion, identical re-ingestion no-op, changed workbook → change
event, D17 graph execution from ingested data, bunker-price change propagation with the
unrelated branch unaffected, Excel provenance in lineage, and rejected ingestion persisted
(status=rejected, no partial publish).

Real .xlsx fixtures built in-memory with openpyxl; real SQLite (and PostgreSQL when
TEST_DATABASE_URL is set); no mocks.
"""

from __future__ import annotations

import io
import json
import os

import openpyxl
import pytest
from sqlalchemy import create_engine, event as sa_event
from sqlalchemy.orm import Session, sessionmaker

from backend.app.config.settings import _normalize_db_url
from backend.app.ingestion import errors as ing_errors
from backend.app.ingestion.excel_connector import read_cells, sha256_hex
from backend.app.ingestion.mappings import PORT_ASSUMPTIONS_V1, get_mapping
from backend.app.ingestion.service import MAX_WORKBOOK_BYTES, IngestionService
from backend.app.persistence.database import Base, ChangeEvent, ExecutionRun, IngestionRun, Result
from backend.app.services.adapter_registry import AdapterRegistry
from backend.app.services.contract_validation import ContractViolationError
from backend.app.services.dataset_values import DatasetValueService
from backend.app.services.orchestration import GraphOrchestrationService
from backend.app.services.results_lineage import ResultsLineageService
from backend.app.ui.port_seed import seed_port_domain

GOLDEN = {
    "bunker_price": 600.0,
    "annual_vessel_calls": 10000.0,
    "average_teu_per_call": 2000.0,
    "average_berth_hours_per_call": 2.0,
    "number_of_berths": 6.0,
    "fuel_burned_in_port_per_call": 2.0,
    "emission_factor": 3.114,
}
# Cell → value in workbook column B, matching PORT_ASSUMPTIONS_V1 (B2..B8).
_ROWS = [
    ("bunker_price", 600),
    ("annual_vessel_calls", 10000),
    ("average_teu_per_call", 2000),
    ("average_berth_hours_per_call", 2.0),
    ("number_of_berths", 6),
    ("fuel_burned_in_port_per_call", 2.0),
    ("emission_factor", 3.114),
]


def make_workbook(values: list | None = None, *, sheet: str = "Assumptions",
                  formula_b2: bool = False) -> bytes:
    """Build a real .xlsx: labels in A2..A8, values in B2..B8."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    rows = values if values is not None else [v for _, v in _ROWS]
    labels = [k for k, _ in _ROWS]
    for i, (label, val) in enumerate(zip(labels, rows)):
        r = i + 2
        ws[f"A{r}"] = label
        ws[f"B{r}"] = val
    if formula_b2:
        ws["B2"] = "=300*2"   # openpyxl writes the formula string; no cached value
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ===========================================================================
# Connector unit tests
# ===========================================================================

class TestConnector:
    def test_reads_all_mapped_cells(self):
        cells = read_cells(make_workbook(), PORT_ASSUMPTIONS_V1)
        assert cells[("Assumptions", "B2")] == 600
        assert cells[("Assumptions", "B8")] == 3.114
        assert len(cells) == 7

    def test_sha256_is_deterministic_by_content(self):
        a, b = make_workbook(), make_workbook()
        assert sha256_hex(a) == sha256_hex(b)
        assert len(sha256_hex(a)) == 64

    def test_sha256_changes_with_content(self):
        base = make_workbook()
        changed = make_workbook([900, 10000, 2000, 2.0, 6, 2.0, 3.114])
        assert sha256_hex(base) != sha256_hex(changed)

    def test_missing_worksheet_raises(self):
        wb = make_workbook(sheet="WrongName")
        with pytest.raises(ing_errors.WorksheetNotFoundError):
            read_cells(wb, PORT_ASSUMPTIONS_V1)

    def test_blank_cell_reads_as_none(self):
        cells = read_cells(make_workbook([None, 10000, 2000, 2.0, 6, 2.0, 3.114]), PORT_ASSUMPTIONS_V1)
        assert cells[("Assumptions", "B2")] is None

    def test_numeric_conversion_preserved(self):
        cells = read_cells(make_workbook([600.5] + [v for _, v in _ROWS[1:]]), PORT_ASSUMPTIONS_V1)
        assert cells[("Assumptions", "B2")] == 600.5

    def test_corrupt_workbook_raises(self):
        with pytest.raises(ing_errors.WorkbookUnreadableError):
            read_cells(b"this is not a zip/xlsx", PORT_ASSUMPTIONS_V1)

    def test_formula_without_cached_value_rejected(self):
        wb = make_workbook(formula_b2=True)
        with pytest.raises(ing_errors.FormulaWithoutCachedValueError):
            read_cells(wb, PORT_ASSUMPTIONS_V1)


# ===========================================================================
# Fixtures
# ===========================================================================

def _sqlite_session() -> Session:
    eng = create_engine("sqlite:///:memory:", connect_args={"check_same_thread": False})

    @sa_event.listens_for(eng, "connect")
    def _fk(conn, _):
        conn.execute("PRAGMA foreign_keys=ON")

    Base.metadata.create_all(bind=eng)
    return sessionmaker(bind=eng, autocommit=False, autoflush=False)()


@pytest.fixture
def db():
    session = _sqlite_session()
    yield session
    session.close()


def _seed_port_for_ingestion(db):
    """Seed the D17 port domain, blank the seeded assumptions value, and clear the seed's
    ChangeEvent so ingestion starts from a clean provenance slate."""
    cfg = seed_port_domain(db)
    ds = DatasetValueService(db).get_dataset(cfg["assumptions_dataset_id"])
    ds.current_value = None
    ds.current_hash = None
    db.query(ChangeEvent).delete()
    db.flush()
    return cfg


@pytest.fixture
def port(db):
    return _seed_port_for_ingestion(db)


def _ingest(db, source_name, values=None, mapping="port_assumptions_v1", **kw):
    svc = IngestionService(db)
    return svc.ingest_excel(source_name=source_name,
                            workbook_bytes=make_workbook(values, **kw), mapping_key=mapping)


def _run_graph(db, cfg):
    return GraphOrchestrationService(db, AdapterRegistry(db)).execute_graph(cfg["terminal_version_id"])


def _summary(db, cfg, go):
    for r in ResultsLineageService(db).list_results_for_run(go.run_id):
        if r.model_version_id == cfg["terminal_version_id"]:
            return json.loads(r.value_json)
    raise AssertionError("no summary result")


# ===========================================================================
# Integration — ingest → dataset → provenance
# ===========================================================================

class TestIngestIntoDataset:
    def test_first_ingestion_writes_dataset(self, db, port):
        res = _ingest(db, "workbook_a.xlsx")
        assert res.status == "ingested" and res.changed is True
        assert res.content_sha256 and len(res.content_sha256) == 64
        assert DatasetValueService(db).get_value(port["assumptions_dataset_id"]) == GOLDEN

    def test_ingestion_run_persisted_with_full_hash(self, db, port):
        res = _ingest(db, "workbook_a.xlsx")
        run = db.get(IngestionRun, res.ingestion_id)
        assert run.status == "ingested"
        assert run.content_sha256 == res.content_sha256          # full 64-char hash persisted
        assert run.mapping_key == "port_assumptions_v1"
        assert port["assumptions_dataset_id"] in json.loads(run.dataset_ids_json)

    def test_change_event_has_excel_provenance(self, db, port):
        res = _ingest(db, "workbook_a.xlsx")
        event = db.get(ChangeEvent, res.change_event_id)
        assert event.source_type == "excel"
        assert event.source_ref.startswith("workbook_a.xlsx@sha256:")
        assert event.source_ref.endswith(res.content_sha256[:12])   # shortened ref, not full hash

    def test_identical_reingestion_is_noop(self, db, port):
        _ingest(db, "workbook_a.xlsx")
        events_before = db.query(ChangeEvent).count()
        res2 = _ingest(db, "workbook_a_copy.xlsx")   # same content, different name
        assert res2.status == "unchanged" and res2.changed is False
        assert res2.graph_run_id is None
        assert db.query(ChangeEvent).count() == events_before

    def test_changed_workbook_creates_new_change_event(self, db, port):
        _ingest(db, "workbook_a.xlsx")
        res2 = _ingest(db, "workbook_b.xlsx", values=[900, 10000, 2000, 2.0, 6, 2.0, 3.114])
        assert res2.status == "ingested" and res2.changed is True
        assert DatasetValueService(db).get_value(port["assumptions_dataset_id"])["bunker_price"] == 900.0


# ===========================================================================
# Integration — D17 graph from ingested data + propagation
# ===========================================================================

class TestGraphFromIngestedData:
    def test_golden_graph_after_ingestion(self, db, port):
        _ingest(db, "workbook_a.xlsx")
        go = _run_graph(db, port)
        assert go.success
        s = _summary(db, port, go)
        assert s["annual_teu"] == pytest.approx(20_000_000.0)
        assert s["annual_fuel_cost_usd"] == pytest.approx(12_000_000.0)
        assert s["fuel_cost_per_teu"] == pytest.approx(0.60)
        assert s["annual_emissions_tco2"] == pytest.approx(62_280.0)
        assert s["emissions_per_teu"] == pytest.approx(0.003114)

    def test_ingestion_triggers_graphrun_and_publishes_summary(self, db, port):
        res = _ingest(db, "workbook_a.xlsx")
        # Ingestion propagates immediately via one GraphRun.
        assert res.graph_run_id is not None
        run = db.get(ExecutionRun, res.graph_run_id)
        assert run.run_kind == "graph" and run.status == "succeeded"
        summary_value = DatasetValueService(db).get_value(port["datasets"]["decision_summary_output"])
        assert summary_value["annual_fuel_cost_usd"] == pytest.approx(12_000_000.0)

    def test_bunker_price_change_propagates_to_cost_not_throughput(self, db, port):
        _ingest(db, "workbook_a.xlsx")
        vals = DatasetValueService(db)
        teu_before = vals.get_value(port["datasets"]["throughput_output"])["annual_teu"]

        res = _ingest(db, "workbook_b.xlsx", values=[900, 10000, 2000, 2.0, 6, 2.0, 3.114])
        assert res.status == "ingested"

        summary = vals.get_value(port["datasets"]["decision_summary_output"])
        # Cost branch reflects the new bunker price...
        assert summary["annual_fuel_cost_usd"] == pytest.approx(20_000.0 * 900.0)
        assert summary["fuel_cost_per_teu"] == pytest.approx(900.0 * 20_000.0 / 20_000_000.0)
        # ...while the throughput (unrelated to bunker price) is unchanged.
        assert vals.get_value(port["datasets"]["throughput_output"])["annual_teu"] == teu_before
        assert summary["annual_teu"] == teu_before

    def test_excel_provenance_reachable_through_lineage(self, db, port):
        _ingest(db, "workbook_a.xlsx")
        go = _run_graph(db, port)
        rl = ResultsLineageService(db)
        by_version = {r.model_version_id: r for r in rl.list_results_for_run(go.run_id)}
        throughput_vid = next(m["version_id"] for m in port["models"] if m["key"] == "throughput")
        edges = rl.list_lineage_to_result(by_version[throughput_vid].id)
        src_edge = next(e for e in edges if e.source_dataset_id == port["assumptions_dataset_id"])
        event = db.get(ChangeEvent, src_edge.source_change_event_id)
        assert event.source_type == "excel"
        assert event.source_ref.startswith("workbook_a.xlsx@sha256:")


# ===========================================================================
# Rejections — persisted, no partial publish
# ===========================================================================

class TestRejections:
    def test_negative_value_rejected_and_persisted(self, db, port):
        res = _ingest(db, "bad.xlsx", values=[-1, 10000, 2000, 2.0, 6, 2.0, 3.114])
        assert res.status == "rejected"
        run = db.get(IngestionRun, res.ingestion_id)
        assert run.status == "rejected" and "below minimum" in (run.error or "")
        # No value published, no change event, no graph run.
        assert DatasetValueService(db).get_dataset(port["assumptions_dataset_id"]).current_value is None
        assert db.query(ChangeEvent).count() == 0
        assert db.query(ExecutionRun).count() == 0

    def test_zero_berths_rejected(self, db, port):
        res = _ingest(db, "bad.xlsx", values=[600, 10000, 2000, 2.0, 0, 2.0, 3.114])
        assert res.status == "rejected"
        assert "number_of_berths" in (db.get(IngestionRun, res.ingestion_id).error or "")

    def test_blank_required_cell_rejected(self, db, port):
        res = _ingest(db, "bad.xlsx", values=[None, 10000, 2000, 2.0, 6, 2.0, 3.114])
        assert res.status == "rejected"
        assert "bunker_price" in (db.get(IngestionRun, res.ingestion_id).error or "")

    def test_wrong_type_rejected(self, db, port):
        res = _ingest(db, "bad.xlsx", values=["六百", 10000, 2000, 2.0, 6, 2.0, 3.114])
        assert res.status == "rejected"
        assert "bunker_price" in (db.get(IngestionRun, res.ingestion_id).error or "")

    def test_missing_worksheet_rejected_and_persisted(self, db, port):
        res = _ingest(db, "bad.xlsx", sheet="Sheet1")
        assert res.status == "rejected"
        run = db.get(IngestionRun, res.ingestion_id)
        assert run.status == "rejected" and "Assumptions" in (run.error or "")
        assert run.content_sha256 is not None    # hash recorded before parsing failed

    def test_corrupt_workbook_rejected(self, db, port):
        res = IngestionService(db).ingest_excel(
            source_name="corrupt.xlsx", workbook_bytes=b"not a workbook",
            mapping_key="port_assumptions_v1")
        assert res.status == "rejected"
        assert db.get(IngestionRun, res.ingestion_id).status == "rejected"

    def test_unsupported_extension_rejected(self, db, port):
        res = IngestionService(db).ingest_excel(
            source_name="data.csv", workbook_bytes=make_workbook(), mapping_key="port_assumptions_v1")
        assert res.status == "rejected"
        assert ".xlsx" in (db.get(IngestionRun, res.ingestion_id).error or "")

    def test_oversized_upload_rejected(self, db, port):
        big = make_workbook() + b"\x00" * (MAX_WORKBOOK_BYTES + 1)
        res = IngestionService(db).ingest_excel(
            source_name="big.xlsx", workbook_bytes=big, mapping_key="port_assumptions_v1")
        assert res.status == "rejected"
        assert "limit" in (db.get(IngestionRun, res.ingestion_id).error or "")

    def test_unknown_mapping_rejected(self, db, port):
        res = IngestionService(db).ingest_excel(
            source_name="a.xlsx", workbook_bytes=make_workbook(), mapping_key="nope")
        assert res.status == "rejected"
        assert "nope" in (db.get(IngestionRun, res.ingestion_id).error or "")

    def test_formula_cell_without_cache_rejected(self, db, port):
        res = _ingest(db, "formula.xlsx", formula_b2=True)
        assert res.status == "rejected"
        assert "formula" in (db.get(IngestionRun, res.ingestion_id).error or "").lower()


# ===========================================================================
# HTTP API
# ===========================================================================

class TestApi:
    @pytest.fixture
    def client(self):
        # TestClient runs endpoints on a worker thread, so a shared-connection StaticPool
        # engine is required for the in-memory DB to be visible across threads.
        from fastapi.testclient import TestClient
        from sqlalchemy.pool import StaticPool
        from backend.app.api.deps import get_db
        from backend.app.main import api

        eng = create_engine("sqlite:///:memory:", connect_args={"check_same_thread": False},
                            poolclass=StaticPool)

        @sa_event.listens_for(eng, "connect")
        def _fk(conn, _):
            conn.execute("PRAGMA foreign_keys=ON")

        Base.metadata.create_all(bind=eng)
        SF = sessionmaker(bind=eng, autocommit=False, autoflush=False)
        seed_db = SF()
        cfg = _seed_port_for_ingestion(seed_db)
        seed_db.commit()
        seed_db.close()

        def _override():
            db = SF()
            try:
                yield db
                db.commit()
            except Exception:
                db.rollback()
                raise
            finally:
                db.close()

        api.dependency_overrides[get_db] = _override
        with TestClient(api, raise_server_exceptions=True) as c:
            yield c, cfg
        api.dependency_overrides.clear()
        eng.dispose()

    def test_upload_ingests(self, client):
        c, port = client
        files = {"file": ("workbook_a.xlsx", make_workbook(),
                          "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        r = c.post("/api/ingestions/excel?mapping=port_assumptions_v1", files=files)
        assert r.status_code == 201
        body = r.json()
        assert body["status"] == "ingested" and body["changed"] is True
        assert len(body["content_sha256"]) == 64

    def test_rejected_upload_is_201_with_status(self, client):
        c, _ = client
        files = {"file": ("bad.xlsx", make_workbook([-1, 10000, 2000, 2.0, 6, 2.0, 3.114]),
                          "application/octet-stream")}
        r = c.post("/api/ingestions/excel?mapping=port_assumptions_v1", files=files)
        assert r.status_code == 201            # durable outcome, not an HTTP error
        assert r.json()["status"] == "rejected"

    def test_unknown_mapping_404(self, client):
        c, _ = client
        files = {"file": ("a.xlsx", make_workbook(), "application/octet-stream")}
        assert c.post("/api/ingestions/excel?mapping=nope", files=files).status_code == 404

    def test_get_provenance_and_list(self, client):
        c, _ = client
        files = {"file": ("workbook_a.xlsx", make_workbook(), "application/octet-stream")}
        ing_id = c.post("/api/ingestions/excel?mapping=port_assumptions_v1", files=files).json()["ingestion_id"]
        got = c.get(f"/api/ingestions/{ing_id}").json()
        assert got["id"] == ing_id and got["status"] == "ingested"
        assert any(r["id"] == ing_id for r in c.get("/api/ingestions").json())

    def test_mappings_listed(self, client):
        c, _ = client
        maps = c.get("/api/ingestions/mappings").json()
        m = next(x for x in maps if x["key"] == "port_assumptions_v1")
        assert len(m["cells"]) == 7


# ===========================================================================
# PostgreSQL (skipped unless TEST_DATABASE_URL is set; rolled back)
# ===========================================================================

_PG_URL = _normalize_db_url(os.environ.get("TEST_DATABASE_URL", ""))


@pytest.mark.skipif(not _PG_URL.startswith("postgresql"), reason="TEST_DATABASE_URL not set")
class TestD18OnPostgreSQL:
    @pytest.fixture()
    def pg_db(self):
        eng = create_engine(_PG_URL, connect_args={"prepare_threshold": None})
        conn = eng.connect()
        trans = conn.begin()
        Base.metadata.create_all(bind=conn)
        session = Session(bind=conn, join_transaction_mode="create_savepoint", autoflush=False)
        yield session
        session.close()
        trans.rollback()
        conn.close()
        eng.dispose()

    def _seed(self, db):
        cfg = seed_port_domain(db)
        ds = DatasetValueService(db).get_dataset(cfg["assumptions_dataset_id"])
        ds.current_value = None
        ds.current_hash = None
        db.flush()
        return cfg

    def test_ingest_and_graph_on_pg(self, pg_db):
        cfg = self._seed(pg_db)
        res = _ingest(pg_db, "workbook_a.xlsx")
        assert res.status == "ingested"
        run = pg_db.get(IngestionRun, res.ingestion_id)
        assert run.content_sha256 == res.content_sha256
        summary = DatasetValueService(pg_db).get_value(cfg["datasets"]["decision_summary_output"])
        assert summary["annual_fuel_cost_usd"] == pytest.approx(12_000_000.0)

    def test_rejected_persisted_on_pg(self, pg_db):
        self._seed(pg_db)
        res = _ingest(pg_db, "bad.xlsx", values=[-1, 10000, 2000, 2.0, 6, 2.0, 3.114])
        assert res.status == "rejected"
        assert pg_db.get(IngestionRun, res.ingestion_id).status == "rejected"
