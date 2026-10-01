"""D26 Stage 2 — authentication, authorization, organization isolation, governed sharing,
HTTP hardening, and migration 007.

The isolation tests are route-coverage tests: they enumerate EVERY API route the running
application serves and attack each one — unauthenticated, and as Northbay with Eastmouth's
identifiers substituted into every path parameter. A new route cannot silently skip the
boundary: it is picked up and tested automatically.
"""

from __future__ import annotations

import hashlib
import io
import json
import re
from datetime import datetime, timedelta, timezone

import openpyxl
import pytest
from fastapi import APIRouter
from fastapi.routing import APIRoute
from fastapi.testclient import TestClient
from sqlalchemy import create_engine, event as sa_event, select, text
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from backend.app import main as main_module
from backend.app.api.deps import get_db
from backend.app.config import settings
from backend.app.main import api
from backend.app.persistence.database import (
    AppUser,
    ApprovedOutput,
    AuditEvent,
    Base,
    ChangeEvent,
    DataContract,
    Dataset,
    ExecutionRun,
    IngestionRun,
    LineageEdge,
    Model,
    ModelVersion,
    Organization,
    Participant,
    Result,
    Scenario,
    ScenarioOverride,
    Baseline,
    UserSession,
)
from backend.app.security import auth as auth_module
from backend.app.security.auth import AuthConfigError, POLICY, effective_auth_mode, required_role
from backend.app.security.passwords import hash_password, verify_password
from backend.app.security.tenancy import TenantViolation, tenant_scope
from backend.app.ui.network_seed import DEMO_PASSWORD, seed_network_demo
from tests.demo_support import db_override

UUID = re.compile(r"[0-9a-f]{8}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{12}")
PUBLIC = {"/api/session", "/api/auth/login", "/api/auth/logout", "/api/build-info",
          # HMAC-authenticated, not session-authenticated. Its refusal is 404 (unknown source)
          # or 401 (bad signature), never based on the caller's user session.
          "/api/webhooks/{source_id}"}


# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------

def _engine():
    eng = create_engine("sqlite:///:memory:", connect_args={"check_same_thread": False},
                        poolclass=StaticPool)

    @sa_event.listens_for(eng, "connect")
    def _fk(conn, _):
        conn.execute("PRAGMA foreign_keys=ON")

    Base.metadata.create_all(bind=eng)
    return eng


@pytest.fixture
def net(monkeypatch):
    eng = _engine()
    SF = sessionmaker(bind=eng)
    db = SF()
    seeded = seed_network_demo(db, extended=False)
    db.commit()
    db.close()
    monkeypatch.setattr(settings, "AUTH_MODE", "required")
    for limiter in (auth_module.login_ip_limiter, auth_module.login_account_limiter, auth_module.write_limiter):
        limiter.reset()
    api.dependency_overrides[get_db] = db_override(SF)
    with TestClient(api) as client:
        yield client, SF, seeded
    api.dependency_overrides.clear()
    eng.dispose()


def login(client: TestClient, email: str, password: str = DEMO_PASSWORD) -> str:
    client.cookies.clear()
    r = client.post("/api/auth/login", json={"email": email, "password": password})
    assert r.status_code == 200, r.text
    return client.cookies.get("sp_csrf")


def api_routes() -> list[tuple[str, str]]:
    """(method, path) for every APIRoute the app serves, however routers are nested."""
    out: set[tuple[str, str]] = set()
    seen: set[int] = set()

    def walk(node):
        if node is None or id(node) in seen:
            return
        seen.add(id(node))
        if isinstance(node, APIRoute):
            for m in node.methods:
                out.add((m, node.path))
        for attr in ("original_router", "router", "app"):
            child = getattr(node, attr, None)
            if child is not None and child is not node:
                walk(child)
        for r in getattr(node, "routes", None) or []:
            walk(r)

    walk(api)
    return sorted(p for p in out if p[1].startswith("/api/"))


def org_ids(SF) -> dict[str, str]:
    db = SF()
    try:
        return {o.org_key: o.id for o in db.execute(select(Organization)).scalars()}
    finally:
        db.close()


def ids_of(SF, org_id: str) -> dict[str, list[str]]:
    """Every identifier one organization owns, grouped by the path parameter names used."""
    db = SF()
    try:
        def q(model):
            return [r.id for r in db.execute(select(model).where(model.organization_id == org_id)).scalars()]
        return {
            "model_id": q(Model), "version_id": q(ModelVersion), "dataset_id": q(Dataset),
            "contract_id": q(DataContract), "run_id": q(ExecutionRun), "result_id": q(Result),
            "baseline_id": q(Baseline), "scenario_id": q(Scenario), "override_id": q(ScenarioOverride),
            "participant_id": q(Participant), "approval_id": q(ApprovedOutput),
            "ingestion_id": q(IngestionRun), "change_event_id": q(ChangeEvent), "edge_id": q(LineageEdge),
        }
    finally:
        db.close()


# ---------------------------------------------------------------------------
# Configuration guards
# ---------------------------------------------------------------------------

class TestAuthMode:
    def test_local_mode_refused_for_a_shared_database(self, monkeypatch):
        monkeypatch.setattr(settings, "AUTH_MODE", "local")
        monkeypatch.setattr(settings, "IS_SQLITE", False)
        with pytest.raises(AuthConfigError):
            effective_auth_mode()

    def test_local_mode_refused_in_production(self, monkeypatch):
        monkeypatch.setattr(settings, "AUTH_MODE", "local")
        monkeypatch.setattr(settings, "IS_PRODUCTION", True)
        with pytest.raises(AuthConfigError):
            effective_auth_mode()

    def test_startup_refuses_an_unsafe_configuration(self, monkeypatch):
        monkeypatch.setattr(settings, "AUTH_MODE", "local")
        monkeypatch.setattr(settings, "IS_SQLITE", False)
        with pytest.raises(AuthConfigError):
            main_module.on_startup()

    def test_unknown_mode_refused(self, monkeypatch):
        monkeypatch.setattr(settings, "AUTH_MODE", "maybe")
        with pytest.raises(AuthConfigError):
            effective_auth_mode()

    def test_required_is_the_default_outside_tests(self):
        src = open(settings.__file__, encoding="utf-8").read()
        assert '"local" if ENVIRONMENT.strip().lower() == "test" else "required"' in src


# ---------------------------------------------------------------------------
# Authentication
# ---------------------------------------------------------------------------

class TestAuthentication:
    def test_every_data_route_requires_sign_in(self, net):
        client, _, _ = net
        client.cookies.clear()
        leaks = []
        for method, path in api_routes():
            if path in PUBLIC:
                continue
            url = re.sub(r"\{[^}]+\}", "x", path).replace(":path", "")
            r = client.request(method, url, json={} if method in ("POST", "PUT", "PATCH") else None)
            if r.status_code != 401:
                leaks.append((method, path, r.status_code))
        assert leaks == []

    def test_session_probe_is_public_and_empty(self, net):
        client, _, _ = net
        client.cookies.clear()
        s = client.get("/api/session").json()
        assert s["authenticated"] is False and s["organizations"] == [] and s["memberships"] == []

    def test_login_sets_hardened_cookies(self, net, monkeypatch):
        client, _, _ = net
        monkeypatch.setattr(settings, "SESSION_COOKIE_SECURE", True)
        r = client.post("/api/auth/login", json={"email": "analyst@northbay.example", "password": DEMO_PASSWORD})
        cookies = r.headers.get_list("set-cookie")
        session = next(c for c in cookies if c.startswith("sp_session="))
        csrf = next(c for c in cookies if c.startswith("sp_csrf="))
        assert "HttpOnly" in session and "SameSite=lax" in session and "Secure" in session
        assert "HttpOnly" not in csrf and "Secure" in csrf

    def test_failed_logins_are_indistinguishable(self, net):
        client, _, _ = net
        a = client.post("/api/auth/login", json={"email": "nobody@nowhere.example", "password": "x" * 14})
        b = client.post("/api/auth/login", json={"email": "analyst@northbay.example", "password": "wrong-password-1"})
        assert a.status_code == b.status_code == 401
        assert a.json() == b.json()

    def test_login_is_throttled(self, net, monkeypatch):
        client, _, _ = net
        for _ in range(settings.LOGIN_MAX_ATTEMPTS):
            client.post("/api/auth/login", json={"email": "viewer@northbay.example", "password": "wrong-password-1"})
        r = client.post("/api/auth/login", json={"email": "viewer@northbay.example", "password": DEMO_PASSWORD})
        assert r.status_code == 429 and int(r.headers["retry-after"]) > 0

    def test_tokens_and_passwords_are_stored_hashed(self, net):
        client, SF, _ = net
        login(client, "analyst@northbay.example")
        token = client.cookies.get("sp_session")
        db = SF()
        rows = db.execute(select(UserSession)).scalars().all()
        assert token not in [r.token_hash for r in rows]
        assert hashlib.sha256(token.encode()).hexdigest() in [r.token_hash for r in rows]
        user = db.execute(select(AppUser).where(AppUser.email == "analyst@northbay.example")).scalar_one()
        assert user.password_hash.startswith("scrypt$") and DEMO_PASSWORD not in user.password_hash
        assert verify_password(DEMO_PASSWORD, user.password_hash)
        db.close()

    def test_absolute_and_idle_expiry(self, net):
        client, SF, _ = net
        login(client, "analyst@northbay.example")
        assert client.get("/api/workspace").status_code == 200
        db = SF()
        row = db.execute(select(UserSession)).scalars().first()
        row.last_seen_at = datetime.now(timezone.utc) - timedelta(minutes=settings.SESSION_IDLE_MINUTES + 5)
        db.commit(); db.close()
        assert client.get("/api/workspace").status_code == 401
        login(client, "analyst@northbay.example")
        db = SF()
        for r in db.execute(select(UserSession).where(UserSession.revoked_at.is_(None))).scalars():
            r.expires_at = datetime.now(timezone.utc) - timedelta(seconds=1)
        db.commit(); db.close()
        assert client.get("/api/workspace").status_code == 401

    def test_logout_revokes_the_session(self, net):
        client, _, _ = net
        csrf = login(client, "analyst@northbay.example")
        token = client.cookies.get("sp_session")
        client.post("/api/auth/logout", headers={"X-CSRF-Token": csrf})
        client.cookies.set("sp_session", token)
        assert client.get("/api/workspace").status_code == 401

    def test_password_change_rules_and_revocation(self, net):
        client, SF, _ = net
        other = TestClient(api)
        other.post("/api/auth/login", json={"email": "analyst@northbay.example", "password": DEMO_PASSWORD})
        csrf = login(client, "analyst@northbay.example")
        h = {"X-CSRF-Token": csrf}
        assert client.post("/api/auth/password", json={"current_password": "nope-nope-nope", "new_password": "x" * 20}, headers=h).status_code == 403
        assert client.post("/api/auth/password", json={"current_password": DEMO_PASSWORD, "new_password": "short"}, headers=h).status_code == 422
        r = client.post("/api/auth/password", json={"current_password": DEMO_PASSWORD, "new_password": "a-new-long-password-2026"}, headers=h)
        assert r.status_code == 200 and r.json()["other_sessions_revoked"] >= 1
        assert other.get("/api/workspace").status_code == 401       # the other session is gone
        assert client.get("/api/workspace").status_code == 200      # this one survives


# ---------------------------------------------------------------------------
# CSRF, origin and write limits
# ---------------------------------------------------------------------------

class TestWriteProtection:
    def _scenario(self, client) -> str:
        return client.get("/api/workspace").json()["scenarios"][0]["id"]

    def test_csrf_token_required_for_writes(self, net):
        client, SF, _ = net
        csrf = login(client, "analyst@northbay.example")
        sid = self._scenario(client)
        assert client.post(f"/api/scenarios/{sid}/execute").status_code == 403
        assert client.post(f"/api/scenarios/{sid}/execute", headers={"X-CSRF-Token": "forged"}).status_code == 403
        assert client.post(f"/api/scenarios/{sid}/execute", headers={"X-CSRF-Token": csrf}).status_code == 201
        db = SF()
        assert db.execute(select(AuditEvent).where(AuditEvent.action == "csrf.rejected")).scalars().first() is not None
        db.close()

    def test_cross_origin_writes_are_refused(self, net):
        client, _, _ = net
        csrf = login(client, "analyst@northbay.example")
        sid = self._scenario(client)
        evil = {"X-CSRF-Token": csrf, "Origin": "https://evil.example"}
        assert client.post(f"/api/scenarios/{sid}/execute", headers=evil).status_code == 403
        ok = {"X-CSRF-Token": csrf, "Origin": "http://testserver"}
        assert client.post(f"/api/scenarios/{sid}/execute", headers=ok).status_code == 201

    def test_write_rate_limit(self, net, monkeypatch):
        client, _, _ = net
        csrf = login(client, "analyst@northbay.example")
        sid = self._scenario(client)
        monkeypatch.setattr(auth_module.write_limiter, "limit", 2)
        auth_module.write_limiter.reset()
        codes = [client.post(f"/api/scenarios/{sid}/execute", headers={"X-CSRF-Token": csrf}).status_code for _ in range(3)]
        assert codes[-1] == 429


# ---------------------------------------------------------------------------
# Roles
# ---------------------------------------------------------------------------

class TestRoles:
    def test_policy_only_names_real_routes(self):
        served = set(api_routes())
        for (method, path) in POLICY:
            assert (method, path) in served, (method, path)

    def test_no_write_route_is_open_to_viewers(self):
        for method, path in api_routes():
            if method in ("POST", "PUT", "PATCH", "DELETE") and path not in PUBLIC:
                assert required_role(method, path) != "viewer", (method, path)

    def test_viewer_reads_but_cannot_run(self, net):
        client, _, _ = net
        csrf = login(client, "viewer@northbay.example")
        ws = client.get("/api/workspace").json()
        sid = ws["scenarios"][0]["id"]
        assert client.post(f"/api/scenarios/{sid}/execute", headers={"X-CSRF-Token": csrf}).status_code == 403

    def test_analyst_runs_but_cannot_approve_or_administer(self, net):
        client, _, _ = net
        csrf = login(client, "analyst@northbay.example")
        h = {"X-CSRF-Token": csrf}
        body = {"participant_id": "p", "dataset_id": "d", "field_name": "f"}
        assert client.post("/api/approved-outputs", json=body, headers=h).status_code == 403
        assert client.post("/api/participants", json={"participant_key": "k", "name": "n"}, headers=h).status_code == 403
        assert client.get("/api/audit").status_code == 403

    def test_approver_approves_admin_administers(self, net):
        client, SF, _ = net
        orgs = org_ids(SF)
        csrf = login(client, "approver@northbay.example")
        h = {"X-CSRF-Token": csrf}
        part = client.get("/api/participants").json()[0]
        summary = next(d for d in client.get("/api/catalog").json() if d["name"] == "decision_summary_output")
        r = client.post("/api/approved-outputs", headers=h, json={
            "participant_id": part["id"], "dataset_id": summary["id"], "field_name": "fuel_cost_per_teu",
            "purpose": "Regional benchmarking", "audience_organization_id": orgs["demo-north-hub"]})
        assert r.status_code == 201, r.text
        assert client.post("/api/participants", json={"participant_key": "k", "name": "n"}, headers=h).status_code == 403
        csrf = login(client, "demo@platform.example")       # admin in Northbay
        assert client.post("/api/participants", json={"participant_key": "northbay-new", "name": "New"},
                           headers={"X-CSRF-Token": csrf}).status_code == 201
        members = client.get(f"/api/organizations/{orgs['port-northbay']}/members").json()
        assert {m["email"] for m in members} >= {"analyst@northbay.example", "demo@platform.example"}


# ---------------------------------------------------------------------------
# Organization isolation
# ---------------------------------------------------------------------------

class TestIsolation:
    def _eastmouth_upload(self, client, SF):
        """Give Eastmouth an ingestion too, so every identifier type exists there."""
        csrf = login(client, "analyst@eastmouth.example")
        raw = client.get("/api/ingestions/template?mapping=port_assumptions_v1").content
        wb = openpyxl.load_workbook(io.BytesIO(raw))
        wb["Assumptions"]["B2"] = 700.0
        buf = io.BytesIO(); wb.save(buf)
        files = {"file": ("east.xlsx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        assert client.post("/api/ingestions/excel?mapping=port_assumptions_v1", files=files,
                           headers={"X-CSRF-Token": csrf}).json()["status"] == "ingested"

    def test_every_route_hides_another_organizations_records(self, net):
        client, SF, _ = net
        self._eastmouth_upload(client, SF)
        orgs = org_ids(SF)
        east = ids_of(SF, orgs["port-eastmouth"])
        assert all(east[k] for k in ("dataset_id", "run_id", "result_id", "scenario_id", "approval_id", "ingestion_id"))
        east_all = {i for v in east.values() for i in v}
        login(client, "analyst@northbay.example")
        leaks = []
        for method, path in api_routes():
            if method != "GET" or path in PUBLIC:
                continue
            params = re.findall(r"\{([^}:]+)", path)
            if not params:
                body = client.get(path).text
                found = east_all & set(UUID.findall(body))
                if found:
                    leaks.append((path, "list", sorted(found)[:2]))
                continue
            pool = next((east.get(p) for p in params if east.get(p)), None)
            if pool is None:
                pool = east["run_id"]
            for ident in pool[:3]:
                url = path
                for p in params:
                    url = re.sub(r"\{" + p + r"[^}]*\}", ident, url)
                if "{" in url:
                    continue
                r = client.get(url)
                if r.status_code == 200 and (east_all & set(UUID.findall(r.text))):
                    leaks.append((path, ident, r.status_code))
        assert leaks == []

    def test_cross_organization_writes_fail_and_change_nothing(self, net):
        client, SF, _ = net
        orgs = org_ids(SF)
        east = ids_of(SF, orgs["port-eastmouth"])
        db = SF()
        before = {d.id: d.current_hash for d in db.execute(select(Dataset).where(Dataset.organization_id == orgs["port-eastmouth"])).scalars()}
        db.close()
        csrf = login(client, "demo@platform.example")        # admin of Northbay, analyst of Eastmouth
        h = {"X-CSRF-Token": csrf, "X-Scope-Org": orgs["port-northbay"]}
        own = client.get("/api/workspace", headers=h).json()
        own_sid = own["scenarios"][0]["id"]
        attempts = [
            client.post(f"/api/scenarios/{east['scenario_id'][0]}/execute", headers=h),
            client.put(f"/api/scenarios/{east['scenario_id'][0]}/overrides", headers=h,
                       json={"overrides": [{"dataset_id": east["dataset_id"][0], "field_name": "bunker_price", "value": 1}]}),
            client.put(f"/api/scenarios/{own_sid}/overrides", headers=h,
                       json={"overrides": [{"dataset_id": east["dataset_id"][0], "field_name": "bunker_price", "value": 1}]}),
            client.post(f"/api/approved-outputs/{east['approval_id'][0]}/revoke", headers=h),
            client.post("/api/changes/propagate", headers=h, json={"dataset_id": east["dataset_id"][0], "value": {"bunker_price": 1}}),
            client.post("/api/graph-executions", headers=h, json={"target_version_id": east["version_id"][0], "dataset_values": {}}),
            client.delete(f"/api/scenarios/{east['scenario_id'][0]}/overrides/{east['override_id'][0]}", headers=h),
            client.post("/api/approved-outputs", headers=h, json={
                "participant_id": client.get("/api/participants", headers=h).json()[0]["id"],
                "dataset_id": east["dataset_id"][0], "field_name": "bunker_price", "purpose": "x"}),
        ]
        assert [a.status_code for a in attempts if a.status_code < 400] == [], [a.status_code for a in attempts]
        db = SF()
        after = {d.id: d.current_hash for d in db.execute(select(Dataset).where(Dataset.organization_id == orgs["port-eastmouth"])).scalars()}
        revoked = db.execute(select(ApprovedOutput).where(ApprovedOutput.id == east["approval_id"][0])).scalar_one()
        db.close()
        assert after == before
        assert revoked.status == "active" or revoked.revocation_reason is not None  # untouched by Northbay

    def test_scope_header_is_checked_against_membership(self, net):
        client, SF, _ = net
        orgs = org_ids(SF)
        login(client, "analyst@northbay.example")
        r = client.get("/api/workspace", headers={"X-Scope-Org": orgs["port-eastmouth"]})
        assert r.status_code == 403
        db = SF()
        denied = db.execute(select(AuditEvent).where(AuditEvent.action == "access.denied")).scalars().all()
        db.close()
        assert any(json.loads(e.detail_json or "{}").get("requested_org") == orgs["port-eastmouth"] for e in denied)

    def test_switching_scope_switches_data_without_leaking(self, net):
        client, SF, _ = net
        orgs = org_ids(SF)
        login(client, "demo@platform.example")
        north = client.get("/api/catalog", headers={"X-Scope-Org": orgs["port-northbay"]}).json()
        east = client.get("/api/catalog", headers={"X-Scope-Org": orgs["port-eastmouth"]}).json()
        assert {d["id"] for d in north}.isdisjoint({d["id"] for d in east})
        teu = lambda cat: next(d for d in cat if d["name"] == "decision_summary_output")["value"]["annual_teu"]
        assert teu(north) == 20_000_000.0 and teu(east) == 11_700_000.0

    def test_orm_refuses_writing_into_another_organization(self, net):
        _, SF, _ = net
        orgs = org_ids(SF)
        db = SF()
        with tenant_scope(db, orgs["port-northbay"]):
            db.add(Dataset(name="sneaky", organization_id=orgs["port-eastmouth"]))
            with pytest.raises(TenantViolation):
                db.flush()
        db.rollback(); db.close()


# ---------------------------------------------------------------------------
# Governed sharing
# ---------------------------------------------------------------------------

class TestGovernedSharing:
    def _setup(self, client, SF, email="approver@northbay.example"):
        csrf = login(client, email)
        h = {"X-CSRF-Token": csrf}
        part = next(p for p in client.get("/api/participants").json() if p["status"] == "active")
        summary = next(d for d in client.get("/api/catalog").json() if d["name"] == "decision_summary_output")
        return h, part, summary

    def test_cannot_share_with_itself_or_an_unknown_audience(self, net):
        client, SF, _ = net
        orgs = org_ids(SF)
        h, part, summary = self._setup(client, SF)
        base = {"participant_id": part["id"], "dataset_id": summary["id"], "field_name": "fuel_cost_per_teu", "purpose": "x"}
        assert client.post("/api/approved-outputs", headers=h, json={**base, "audience_organization_id": orgs["port-northbay"]}).status_code == 422
        assert client.post("/api/approved-outputs", headers=h, json={**base, "audience_organization_id": "nope"}).status_code == 422
        db = SF()
        assert db.execute(select(AuditEvent).where(AuditEvent.action == "approval.denied")).scalars().first() is not None
        db.close()

    def test_grant_and_revoke_are_attributed_and_audited(self, net):
        client, SF, _ = net
        orgs = org_ids(SF)
        h, part, summary = self._setup(client, SF)
        a = client.post("/api/approved-outputs", headers=h, json={
            "participant_id": part["id"], "dataset_id": summary["id"], "field_name": "fuel_cost_per_teu",
            "purpose": "Regional benchmarking", "audience_organization_id": orgs["demo-north-hub"]}).json()
        assert a["approved_by_user_id"] and a["audience_organization_id"] == orgs["demo-north-hub"]
        r = client.post(f"/api/approved-outputs/{a['id']}/revoke", headers=h, json={"reason": "Methodology changed"}).json()
        assert r["status"] == "revoked" and r["revoked_by_user_id"] == a["approved_by_user_id"]
        assert r["revocation_reason"] == "Methodology changed"
        actions = [e["action"] for e in client.get("/api/audit").json()]
        assert "approval.created" in actions and "approval.revoked" in actions

    def test_scenario_results_are_never_shared_implicitly(self, net):
        client, SF, _ = net
        h, part, summary = self._setup(client, SF, "demo@platform.example")
        sid = client.get("/api/workspace").json()["scenarios"][0]["id"]
        before = {(x["field_name"], x["value"]) for x in client.get("/api/exposed-outputs").json()}
        assert client.post(f"/api/scenarios/{sid}/execute", headers=h).status_code == 201
        after = {(x["field_name"], x["value"]) for x in client.get("/api/exposed-outputs").json()}
        assert before == after
        assert all(x["value_source"] == "published" for x in client.get("/api/exposed-outputs").json())

    def test_an_explicitly_shared_run_result_exposes_that_run_only(self, net):
        client, SF, _ = net
        orgs = org_ids(SF)
        h, part, summary = self._setup(client, SF, "demo@platform.example")
        ws = client.get("/api/workspace").json()
        run_id = ws["scenarios"][0]["run"]["id"]
        a = client.post("/api/approved-outputs", headers=h, json={
            "participant_id": part["id"], "dataset_id": summary["id"], "field_name": "annual_fuel_cost_usd",
            "purpose": "Joint review of a higher bunker price", "audience_organization_id": orgs["demo-north-hub"],
            "source_run_id": run_id})
        assert a.status_code == 201, a.text
        shared = next(x for x in client.get("/api/exposed-outputs").json()
                      if x["field_name"] == "annual_fuel_cost_usd" and x["source_run_id"] == run_id)
        assert shared["value_source"] == "run" and shared["value"] == 15_000_000.0
        catalog = client.get("/api/catalog").json()
        published = next(d for d in catalog if d["name"] == "decision_summary_output")["value"]["annual_fuel_cost_usd"]
        assert published == 12_000_000.0                         # the dataset itself is untouched

    def test_another_organizations_run_cannot_be_shared(self, net):
        client, SF, _ = net
        orgs = org_ids(SF)
        east_run = ids_of(SF, orgs["port-eastmouth"])["run_id"][0]
        h, part, summary = self._setup(client, SF)
        r = client.post("/api/approved-outputs", headers=h, json={
            "participant_id": part["id"], "dataset_id": summary["id"], "field_name": "annual_teu",
            "purpose": "x", "audience_organization_id": orgs["demo-north-hub"], "source_run_id": east_run})
        assert r.status_code in (409, 422)

    def test_revoked_and_expired_approvals_are_not_exposed(self, net):
        client, SF, _ = net
        login(client, "approver@northbay.example")
        exposed = {x["field_name"] for x in client.get("/api/exposed-outputs").json()}
        assert "annual_fuel_cost_usd" not in exposed            # seeded as revoked
        db = SF()
        a = db.execute(select(ApprovedOutput).where(ApprovedOutput.field_name == "annual_teu",
                                                    ApprovedOutput.status == "active")).scalars().first()
        a.expires_at = datetime.now(timezone.utc) - timedelta(minutes=1)
        org = a.organization_id
        db.commit(); db.close()
        client.cookies.clear()
        login(client, "approver@northbay.example" if org == org_ids(SF)["port-northbay"] else "analyst@eastmouth.example")
        assert "annual_teu" not in {x["field_name"] for x in client.get("/api/exposed-outputs").json()}


# ---------------------------------------------------------------------------
# HTTP hardening
# ---------------------------------------------------------------------------

class TestHttpHardening:
    def test_security_headers_per_surface(self):
        with TestClient(api) as c:
            app_csp = c.get("/app/").headers.get("content-security-policy", "")
            assert "script-src 'self'" in app_csp and "unsafe-inline" not in app_csp.split("style-src")[0]
            ui = c.get("/ui").headers
            assert "frame-ancestors 'none'" in ui["content-security-policy"]
            h = c.get("/api/build-info").headers
            assert h["x-content-type-options"] == "nosniff"
            assert h["x-frame-options"] == "DENY"
            assert h["cache-control"] == "no-store"
            assert h["referrer-policy"] == "strict-origin-when-cross-origin"
            assert re.fullmatch(r"[0-9a-f]{32}", h["x-request-id"])

    def test_request_id_is_echoed_when_sane(self):
        with TestClient(api) as c:
            assert c.get("/health", headers={"X-Request-ID": "trace-12345678"}).headers["x-request-id"] == "trace-12345678"
            assert c.get("/health", headers={"X-Request-ID": "bad id!"}).headers["x-request-id"] != "bad id!"

    def test_hsts_only_when_enabled(self, monkeypatch):
        with TestClient(api) as c:
            assert "strict-transport-security" not in c.get("/health").headers
            monkeypatch.setattr(settings, "HSTS_ENABLED", True)
            assert "max-age=31536000" in c.get("/health").headers["strict-transport-security"]

    def test_oversized_bodies_are_refused(self, monkeypatch):
        monkeypatch.setattr(settings, "MAX_REQUEST_BYTES", 1024)
        with TestClient(api) as c:
            r = c.post("/api/auth/login", content=b"x" * 5000, headers={"content-type": "application/json"})
            assert r.status_code == 413

    def test_unhandled_errors_never_leak_details(self):
        router = APIRouter()

        @router.get("/api/_boom_for_test")
        def boom():
            raise RuntimeError("database password is hunter2")

        api.include_router(router)
        try:
            with TestClient(api, raise_server_exceptions=False) as c:
                r = c.get("/api/_boom_for_test")
            assert r.status_code == 500
            assert "hunter2" not in r.text and r.json()["detail"] == "Internal server error."
            assert r.json()["request_id"]
        finally:
            api.router.routes[:] = [rt for rt in api.router.routes if getattr(rt, "path", "") != "/api/_boom_for_test"]

    def test_legacy_pages_carry_the_compatibility_shim(self):
        with TestClient(api) as c:
            for page in ("/ui", "/ui/start", "/ui/governance"):
                assert "D26 compatibility shim" in c.get(page).text


# ---------------------------------------------------------------------------
# Migration 007 and the network seed
# ---------------------------------------------------------------------------

class TestMigration007:
    def _cfg(self, url):
        from alembic.config import Config
        cfg = Config("alembic.ini")
        cfg.attributes["sqlalchemy_url"] = url
        return cfg

    def test_full_lifecycle_and_backfill(self, tmp_path):
        from alembic import command
        url = f"sqlite:///{tmp_path / 'lc.db'}"
        cfg = self._cfg(url)
        command.upgrade(cfg, "head")
        command.downgrade(cfg, "base")
        command.upgrade(cfg, "006")
        eng = create_engine(url)
        with eng.begin() as c:
            c.execute(text("INSERT INTO dataset (id, name, created_at, updated_at) VALUES ('d1', 'x', CURRENT_TIMESTAMP, CURRENT_TIMESTAMP)"))
        command.upgrade(cfg, "head")
        with eng.begin() as c:
            assert c.execute(text("SELECT org_key FROM organization")).scalars().all() == ["default"]
            assert c.execute(text("SELECT organization_id IS NOT NULL FROM dataset WHERE id='d1'")).scalar() == 1
        command.downgrade(cfg, "006")               # no cross-org duplicates yet: clean rollback
        command.upgrade(cfg, "head")
        eng.dispose()

    def test_downgrade_refuses_to_merge_organizations(self, tmp_path):
        from alembic import command
        url = f"sqlite:///{tmp_path / 'dup.db'}"
        cfg = self._cfg(url)
        command.upgrade(cfg, "head")
        eng = create_engine(url)
        with eng.begin() as c:
            for oid in ("o1", "o2"):
                c.execute(text("INSERT INTO organization (id, org_key, name, kind, status) VALUES (:i, :i, :i, 'port', 'active')"), {"i": oid})
                c.execute(text("INSERT INTO dataset (id, name, organization_id, created_at, updated_at) VALUES (:d, 'port_assumptions', :o, CURRENT_TIMESTAMP, CURRENT_TIMESTAMP)"), {"d": f"d-{oid}", "o": oid})
        with pytest.raises(RuntimeError, match="Cannot downgrade 007"):
            command.downgrade(cfg, "006")
        eng.dispose()


class TestNetworkSeed:
    def test_hierarchy_and_private_federations(self, net):
        _, SF, seeded = net
        db = SF()
        orgs = {o.org_key: o for o in db.execute(select(Organization)).scalars()}
        assert orgs["port-northbay"].parent_id == orgs["demo-north-hub"].id
        assert orgs["demo-north-hub"].parent_id == orgs["demo-national"].id
        assert orgs["demo-national"].parent_id == orgs["demo-network"].id
        for key in ("port-northbay", "port-eastmouth", "port-southreach"):
            n = db.execute(select(Dataset).where(Dataset.organization_id == orgs[key].id,
                                                 Dataset.name == "port_assumptions")).scalars().all()
            assert len(n) == 1
        assert all(json.loads(o.meta_json)["synthetic"] for o in orgs.values())
        db.close()

    def test_idempotent(self, net):
        _, SF, _ = net
        db = SF()
        again = seed_network_demo(db, extended=False)
        db.commit(); db.close()
        assert again["approvals_created"] == 0 and again["accounts_created"] == 0

    def test_never_runs_against_a_non_sqlite_database(self):
        class FakeBind:
            class dialect:
                name = "postgresql"

        class FakeSession:
            def get_bind(self):
                return FakeBind()

        assert seed_network_demo(FakeSession()) == {}
