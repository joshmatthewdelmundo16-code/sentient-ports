"""D24 — parity, live productization and the guided Excel workflow.

Covers the four things D24 set out to fix:

  1. `/ui/governance` returning 404 in a real runtime, and the *reason* it could happen
     invisibly (nothing reported which build was answering).
  2. A first-time user having no in-product explanation of the system.
  3. The Excel lifecycle being invisible and commit-only — no template, no dry run.
  4. Startup writing schema and demo rows into whatever database it was pointed at.

Real SQLite throughout, no mocks. Every business number asserted here is read back from
the API rather than written into the test.
"""

from __future__ import annotations

import io
import os

import openpyxl
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine, event as sa_event
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

os.environ.setdefault("DATABASE_URL", "sqlite:///:memory:")
os.environ.setdefault("STORAGE_ROOT", "storage_test_d24_tmp")

from backend.app import buildinfo  # noqa: E402
from backend.app.api.deps import get_db  # noqa: E402
from backend.app.main import api  # noqa: E402
from backend.app.persistence.database import Base  # noqa: E402
from backend.app.ui import decision_state, demo_state  # noqa: E402
from backend.app.ui.governance_seed import seed_governance_demo  # noqa: E402
from backend.app.ui.port_decision_seed import seed_port_decision  # noqa: E402

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MAPPING = "port_assumptions_v1"


# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------

def _engine():
    eng = create_engine("sqlite:///:memory:", connect_args={"check_same_thread": False},
                        poolclass=StaticPool)

    @sa_event.listens_for(eng, "connect")
    def _fk(conn, _):
        conn.execute("PRAGMA foreign_keys=ON")

    Base.metadata.create_all(bind=eng)
    return eng


def _db_override(SF):
    def _get_db():
        s = SF()
        try:
            yield s
            s.commit()
        except Exception:
            s.rollback()
            raise
        finally:
            s.close()
    return _get_db


@pytest.fixture
def client():
    """A fully seeded app: port domain, baseline/scenario, and governance demo."""
    eng = _engine()
    SF = sessionmaker(bind=eng)
    db = SF()
    cfg = seed_port_decision(db)
    seed_governance_demo(db)
    db.commit()
    db.close()
    api.dependency_overrides[get_db] = _db_override(SF)
    with TestClient(api, raise_server_exceptions=True) as c:
        orig_dec, orig_demo = decision_state.config, demo_state.config
        decision_state.config = cfg
        try:
            yield c
        finally:
            decision_state.config, demo_state.config = orig_dec, orig_demo
            api.dependency_overrides.clear()


@pytest.fixture
def session_factory():
    eng = _engine()
    return sessionmaker(bind=eng)


def _template_bytes(client) -> bytes:
    r = client.get(f"/api/ingestions/template?mapping={MAPPING}")
    assert r.status_code == 200
    return r.content


def _edited(template: bytes, cell: str, value) -> bytes:
    wb = openpyxl.load_workbook(io.BytesIO(template))
    wb["Assumptions"][cell] = value
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _upload(client, path: str, data: bytes, name: str = "port.xlsx"):
    return client.post(f"{path}?mapping={MAPPING}", files={"file": (name, data, XLSX)})


# ---------------------------------------------------------------------------
# 1 — The reported defect: /ui/governance
# ---------------------------------------------------------------------------

class TestGovernanceRoute:
    """The 404 the user saw. The route must exist and be reachable from the product."""

    def test_ui_governance_is_served(self, client):
        r = client.get("/ui/governance")
        assert r.status_code == 200
        assert "Federation governance" in r.text

    def test_ui_governance_has_no_trailing_slash_trap(self, client):
        # A redirect is fine; a 404 is not.
        r = client.get("/ui/governance/", follow_redirects=True)
        assert r.status_code in (200, 307, 308)

    def test_governance_is_linked_from_the_main_navigation(self, client):
        """The page existed in D23 but was only reachable from inside a collapsed panel."""
        nav = client.get("/ui").text.split("</nav>")[0]
        assert 'href="/ui/governance"' in nav

    def test_governance_apis_the_page_calls_all_answer(self, client):
        for path in ("/api/governance/summary", "/api/participants",
                     "/api/approved-outputs", "/api/exposed-outputs"):
            assert client.get(path).status_code == 200, path


class TestBuildSelfReport:
    """Why the 404 was invisible: nothing said which build was answering."""

    def test_build_info_is_served(self, client):
        info = client.get("/api/build-info").json()
        # The build identifies itself; D25–D27 bumped both (the assertion follows the setting,
        # so a future phase cannot silently ship with a stale label).
        from backend.app.config.settings import APP_PHASE, APP_VERSION
        assert info["phase"] == APP_PHASE
        assert info["version"] == APP_VERSION

    def test_every_declared_capability_is_actually_served(self, client):
        info = client.get("/api/build-info").json()
        assert info["missing_capabilities"] == []
        assert all(info["capabilities"].values())

    def test_capabilities_are_derived_from_the_real_route_table(self, client):
        """Not a hand-maintained list — it must reflect routes that genuinely exist."""
        paths = buildinfo.route_paths(api)
        for capability, path in buildinfo.CAPABILITY_ROUTES.items():
            assert path in paths, f"{capability} claims {path}, which is not routed"

    def test_route_discovery_survives_the_wrapped_include_shape(self, client):
        """FastAPI 0.141 stops flattening include_router; a flat read sees ~7 routes."""
        assert len(buildinfo.route_paths(api)) > 40

    def test_a_build_missing_a_route_is_reported_as_missing(self):
        """Simulate the stale server: a build without the governance UI must say so."""
        class FakeRoute:
            def __init__(self, path): self.path = path

        class FakeApp:
            routes = [FakeRoute("/ui"), FakeRoute("/health")]

        caps = buildinfo.capabilities(FakeApp())
        assert caps["ui.workspace"] is True
        assert caps["ui.governance"] is False
        assert "ui.governance" in buildinfo.missing_capabilities(FakeApp())

    def test_health_reports_missing_capabilities(self, client):
        assert client.get("/health").json()["missing_capabilities"] == []

    def test_build_info_leaks_no_credentials(self, client):
        info = client.get("/api/build-info").json()
        assert info["database"] in ("sqlite", "postgresql")
        # Database family only — never a URL, host, user or password.
        assert "://" not in str(info)
        assert "@" not in str(info)


# ---------------------------------------------------------------------------
# 2 — Onboarding
# ---------------------------------------------------------------------------

class TestOnboarding:
    def test_start_here_page_is_served(self, client):
        r = client.get("/ui/start")
        assert r.status_code == 200
        assert "Start here" in r.text

    def test_root_opens_the_product_not_json(self, client):
        """Before D24, / redirected to /health — a JSON blob as the first impression.

        D25: the React product (/app) is now the primary experience; the legacy Start Here
        page stays reachable at /ui/start and via PRIMARY_UI=legacy (see test_d25_product).
        """
        r = client.get("/", follow_redirects=False)
        assert r.status_code in (307, 302)
        assert r.headers["location"] == "/app/"

    def test_start_here_is_linked_from_the_workspace(self, client):
        assert 'href="/ui/start"' in client.get("/ui").text

    @pytest.mark.parametrize("topic", [
        "baseline",          # what the baseline represents
        "scenario",          # what a scenario represents
        "bunker_price",      # which assumptions can be changed
        "Download the workbook",   # how to import the workbook
        "dry run",           # what validation does
        "Commit",            # what happens after
        "propagate",         # how the change reaches dependent models
        "provenance",        # where provenance is shown
        "Governance",        # where governance is shown
        "synthetic",         # honesty about the data
    ])
    def test_start_here_covers_the_required_topics(self, client, topic):
        # Collapse whitespace: HTML wraps phrases across lines, and that is not a gap.
        text = " ".join(client.get("/ui/start").text.lower().split())
        assert topic.lower() in text

    def test_start_here_declares_what_is_not_implemented(self, client):
        """Missing public capabilities must be labelled, not hidden."""
        text = client.get("/ui/start").text
        for absent in ("Operational optimization", "Dynamic master planning",
                       "Live sensor feeds"):
            assert absent in text
        assert "Not implemented" in text

    def test_workspace_exposes_the_excel_workflow(self, client):
        html = client.get("/ui").text
        assert 'id="excel-section"' in html
        assert 'href="#excel-section"' in html


# ---------------------------------------------------------------------------
# 3 — Guided Excel workflow
# ---------------------------------------------------------------------------

class TestMappingPreview:
    def test_mapping_preview_names_cell_field_unit_and_current_value(self, client):
        p = client.get(f"/api/ingestions/mappings/{MAPPING}/preview").json()
        fields = {f["field"]: f for f in p["fields"]}
        bunker = fields["bunker_price"]
        assert (bunker["worksheet"], bunker["cell"]) == ("Assumptions", "B2")
        assert bunker["unit"] == "USD/t"
        assert bunker["current_value"] is not None
        assert p["notes"], "connector limitations must be stated"

    def test_unknown_mapping_is_a_404(self, client):
        assert client.get("/api/ingestions/mappings/nope/preview").status_code == 404


class TestTemplate:
    def test_template_is_a_real_xlsx_download(self, client):
        r = client.get(f"/api/ingestions/template?mapping={MAPPING}")
        assert r.status_code == 200
        assert r.headers["content-type"] == XLSX
        assert "attachment" in r.headers["content-disposition"]
        assert r.content[:2] == b"PK"  # zip container

    def test_template_is_prefilled_with_the_current_values(self, client):
        """A blank form would make the first upload look like seven changes."""
        current = {
            f["field"]: f["current_value"]
            for f in client.get(f"/api/ingestions/mappings/{MAPPING}/preview").json()["fields"]
        }
        wb = openpyxl.load_workbook(io.BytesIO(_template_bytes(client)))
        assert wb["Assumptions"]["B2"].value == current["bunker_price"]
        assert wb["Assumptions"]["B6"].value == current["number_of_berths"]

    def test_template_documents_the_cell_map_and_units(self, client):
        wb = openpyxl.load_workbook(io.BytesIO(_template_bytes(client)))
        assert "How to use" in wb.sheetnames
        assert wb["Assumptions"]["C2"].value == "USD/t"
        assert wb["Assumptions"]["D2"].value == "bunker_price"

    def test_template_round_trips_through_the_connector_unedited(self, client):
        """Downloading and re-uploading without editing must be a no-op, not a change."""
        preview = _upload(client, "/api/ingestions/excel/validate",
                          _template_bytes(client)).json()
        assert preview["valid"] is True
        assert preview["would_change"] is False
        assert preview["changed_fields"] == []

    def test_unknown_mapping_template_is_a_404(self, client):
        assert client.get("/api/ingestions/template?mapping=nope").status_code == 404


class TestValidateIsADryRun:
    def test_validate_reports_the_changed_field_only(self, client):
        data = _edited(_template_bytes(client), "B2", 900.0)
        p = _upload(client, "/api/ingestions/excel/validate", data).json()
        assert p["valid"] is True
        assert p["would_change"] is True
        assert p["changed_fields"] == ["bunker_price"]
        changed = [f for f in p["fields"] if f["changed"]]
        assert len(changed) == 1
        assert changed[0]["new_value"] == 900.0
        assert changed[0]["current_value"] != 900.0
        assert changed[0]["unit"] == "USD/t"

    def test_validate_writes_nothing(self, client):
        before_ingestions = len(client.get("/api/ingestions").json())
        before_runs = len(client.get("/api/executions").json())
        before = client.get(f"/api/ingestions/mappings/{MAPPING}/preview").json()

        _upload(client, "/api/ingestions/excel/validate",
                _edited(_template_bytes(client), "B2", 1234.0))

        assert len(client.get("/api/ingestions").json()) == before_ingestions
        assert len(client.get("/api/executions").json()) == before_runs
        assert client.get(f"/api/ingestions/mappings/{MAPPING}/preview").json() == before

    def test_validate_rejects_a_contract_violation_before_any_write(self, client):
        """number_of_berths has a contract minimum of 1 — a port cannot have zero berths."""
        data = _edited(_template_bytes(client), "B6", 0)
        p = _upload(client, "/api/ingestions/excel/validate", data).json()
        assert p["valid"] is False
        bad = [f for f in p["fields"] if f["field"] == "number_of_berths"][0]
        assert bad["valid"] is False
        assert bad["message"]
        assert len(client.get("/api/ingestions").json()) == 0

    def test_validate_reports_an_unreadable_workbook_without_a_500(self, client):
        r = _upload(client, "/api/ingestions/excel/validate", b"this is not a workbook")
        assert r.status_code == 200
        body = r.json()
        assert body["valid"] is False
        assert body["errors"]

    def test_validate_rejects_a_non_xlsx_name(self, client):
        r = _upload(client, "/api/ingestions/excel/validate", b"x,y\n1,2\n", name="data.csv")
        assert r.status_code == 200
        assert r.json()["valid"] is False
        assert any(".xlsx" in e for e in r.json()["errors"])

    def test_validate_reports_the_content_hash(self, client):
        p = _upload(client, "/api/ingestions/excel/validate", _template_bytes(client)).json()
        assert len(p["content_sha256"]) == 64

    def test_validate_states_the_connector_limitations(self, client):
        notes = " ".join(
            _upload(client, "/api/ingestions/excel/validate", _template_bytes(client))
            .json()["notes"]
        ).lower()
        assert "formula" in notes and "never evaluated" in notes
        assert "macro" in notes

    def test_unknown_mapping_validate_is_a_404(self, client):
        r = client.post("/api/ingestions/excel/validate?mapping=nope",
                        files={"file": ("a.xlsx", b"PK", XLSX)})
        assert r.status_code == 404


class TestValidateAgreesWithCommit:
    def test_a_valid_preview_is_followed_by_a_successful_commit(self, client):
        data = _edited(_template_bytes(client), "B2", 900.0)
        preview = _upload(client, "/api/ingestions/excel/validate", data).json()
        assert preview["valid"] is True

        commit = _upload(client, "/api/ingestions/excel", data).json()
        assert commit["status"] == "ingested"
        assert commit["changed"] is True
        assert commit["content_sha256"] == preview["content_sha256"]
        assert commit["error"] is None

    def test_an_invalid_preview_is_followed_by_a_rejected_commit(self, client):
        data = _edited(_template_bytes(client), "B6", 0)
        assert _upload(client, "/api/ingestions/excel/validate", data).json()["valid"] is False

        commit = _upload(client, "/api/ingestions/excel", data).json()
        assert commit["status"] == "rejected"
        assert commit["error"]


class TestExcelToFederationChain:
    """Stage G — one edited cell, traced all the way to a changed KPI."""

    def test_the_whole_chain_runs_from_one_edited_cell(self, client):
        datasets = {d["name"]: d["id"] for d in client.get("/api/datasets").json()}
        before = client.get(f"/api/datasets/{datasets['port_fuel_cost_output']}").json()
        baseline_cost = before["current_value"]["annual_fuel_cost_usd"]

        data = _edited(_template_bytes(client), "B2", 900.0)
        r = _upload(client, "/api/ingestions/excel", data).json()

        # ingestion record → change event → graph run
        assert r["ingestion_id"]
        assert r["change_event_id"]
        assert r["graph_run_id"]

        # provenance is retrievable
        run = client.get(f"/api/ingestions/{r['ingestion_id']}").json()
        assert run["content_sha256"] == r["content_sha256"]
        assert run["status"] == "ingested"

        # dependent model actually re-ran and published a different value
        after = client.get(f"/api/datasets/{datasets['port_fuel_cost_output']}").json()
        new_cost = after["current_value"]["annual_fuel_cost_usd"]
        assert new_cost != baseline_cost

        # the terminal fan-in model reflects it too
        summary = client.get(f"/api/datasets/{datasets['decision_summary_output']}").json()
        assert summary["current_value"]["annual_fuel_cost_usd"] == new_cost

        # results are attributed to the run the ingestion triggered
        results = client.get(f"/api/executions/{r['graph_run_id']}/results").json()
        assert len(results) >= 5

    def test_only_dependent_outputs_move(self, client):
        """Emissions do not depend on price; the change must not touch them."""
        datasets = {d["name"]: d["id"] for d in client.get("/api/datasets").json()}
        before = client.get(f"/api/datasets/{datasets['port_emissions_output']}").json()

        _upload(client, "/api/ingestions/excel", _edited(_template_bytes(client), "B2", 900.0))

        after = client.get(f"/api/datasets/{datasets['port_emissions_output']}").json()
        assert after["current_value"] == before["current_value"]

    def test_recommitting_the_same_workbook_changes_nothing(self, client):
        data = _edited(_template_bytes(client), "B2", 900.0)
        first = _upload(client, "/api/ingestions/excel", data).json()
        assert first["status"] == "ingested"

        second = _upload(client, "/api/ingestions/excel", data).json()
        assert second["status"] == "unchanged"
        assert second["changed"] is False
        assert second["graph_run_id"] is None
        # Still recorded — an unchanged upload is provenance, not a non-event.
        assert second["ingestion_id"] != first["ingestion_id"]

    def test_the_ingestion_is_listed_as_a_source(self, client):
        _upload(client, "/api/ingestions/excel", _edited(_template_bytes(client), "B2", 900.0))
        ingestions = client.get("/api/ingestions").json()
        assert ingestions
        assert ingestions[0]["source_name"] == "port.xlsx"
        assert ingestions[0]["mapping_key"] == MAPPING


# ---------------------------------------------------------------------------
# 4 — Governance visibility
# ---------------------------------------------------------------------------

class TestGovernanceDemo:
    def test_participants_are_seeded_so_the_page_is_not_blank(self, client):
        assert len(client.get("/api/participants").json()) >= 4

    def test_every_seeded_record_is_labelled_synthetic(self, client):
        """Nothing here may read as a real data-sharing agreement."""
        for p in client.get("/api/participants").json():
            assert p["metadata"]["synthetic"] is True
            assert "not a real" in p["metadata"]["disclaimer"].lower()
        for a in client.get("/api/approved-outputs").json():
            assert a["metadata"]["synthetic"] is True

    def test_exposed_outputs_exclude_inactive_participants(self, client):
        exposed = client.get("/api/exposed-outputs").json()
        assert exposed, "the boundary should not be empty"
        assert all(e["participant_key"] != "prospective-partner" for e in exposed)

    def test_exposed_outputs_exclude_revoked_approvals(self, client):
        approvals = client.get("/api/approved-outputs").json()
        revoked = [a for a in approvals if a["status"] == "revoked"]
        assert revoked, "the demo must include a revoked approval"
        exposed = {(e["field_name"]) for e in client.get("/api/exposed-outputs").json()}
        for a in revoked:
            assert a["field_name"] not in exposed

    def test_exposed_values_are_real_computed_values(self, client):
        """The boundary reports values the models produced, not placeholders."""
        datasets = {d["id"]: d["name"] for d in client.get("/api/datasets").json()}
        for e in client.get("/api/exposed-outputs").json():
            assert e["value_present"] is True
            assert datasets[e["dataset_id"]] == e["dataset_name"]

    def test_governance_page_labels_synthetic_data(self, client):
        assert "synthetic-note" in client.get("/ui/governance").text

    def test_seed_is_idempotent(self, session_factory):
        db = session_factory()
        seed_port_decision(db)
        first = seed_governance_demo(db)
        db.commit()
        second = seed_governance_demo(db)
        db.commit()
        third = seed_governance_demo(db)
        db.commit()
        assert first["participants"] == second["participants"] == third["participants"]
        assert second["approvals"] == 0 and third["approvals"] == 0
        db.close()

    def test_seed_refuses_a_non_sqlite_database(self):
        """A shared PostgreSQL must never gain demo governance rows from a startup."""
        class FakeBind:
            class dialect: name = "postgresql"

        class FakeSession:
            def get_bind(self): return FakeBind()

        assert seed_governance_demo(FakeSession()) == {}


# ---------------------------------------------------------------------------
# 5 — Startup safety
# ---------------------------------------------------------------------------

class TestStartupSafety:
    def test_sqlite_keeps_the_zero_install_developer_defaults(self):
        from backend.app.config import settings
        assert settings.DATABASE_URL.startswith("sqlite")
        assert settings.AUTO_CREATE_SCHEMA is True
        assert settings.DEMO_SEED_ENABLED is True

    def test_the_gates_default_off_for_a_shared_database(self, monkeypatch):
        """Re-import settings as if DATABASE_URL pointed at PostgreSQL."""
        import importlib
        from backend.app.config import settings as settings_module

        monkeypatch.setenv("DATABASE_URL", "postgresql://u:p@example.invalid:5432/db")
        monkeypatch.delenv("AUTO_CREATE_SCHEMA", raising=False)
        monkeypatch.delenv("DEMO_SEED_ENABLED", raising=False)
        reloaded = importlib.reload(settings_module)
        try:
            assert reloaded.AUTO_CREATE_SCHEMA is False
            assert reloaded.DEMO_SEED_ENABLED is False
        finally:
            monkeypatch.undo()
            importlib.reload(settings_module)

    def test_an_operator_can_still_opt_in_explicitly(self, monkeypatch):
        import importlib
        from backend.app.config import settings as settings_module

        monkeypatch.setenv("DATABASE_URL", "postgresql://u:p@example.invalid:5432/db")
        monkeypatch.setenv("AUTO_CREATE_SCHEMA", "true")
        reloaded = importlib.reload(settings_module)
        try:
            assert reloaded.AUTO_CREATE_SCHEMA is True
        finally:
            monkeypatch.undo()
            importlib.reload(settings_module)

    def test_public_base_url_is_not_invented(self):
        """The UI may only claim a public address an operator actually configured."""
        from backend.app.config import settings
        assert settings.PUBLIC_BASE_URL == ""


# ---------------------------------------------------------------------------
# 6 — Deployment configuration
# ---------------------------------------------------------------------------

class TestDeploymentArtifacts:
    @pytest.fixture(scope="class")
    def platform_dir(self):
        from pathlib import Path
        return Path(__file__).resolve().parents[1]

    def test_requirements_is_utf8_and_pip_readable(self, platform_dir):
        """It was UTF-16, which `pip install -r` cannot parse — no container could build."""
        raw = (platform_dir / "requirements.txt").read_bytes()
        assert not raw.startswith(b"\xff\xfe") and not raw.startswith(b"\xfe\xff")
        assert b"\x00" not in raw
        text = raw.decode("utf-8")
        packages = [l for l in text.splitlines() if l.strip() and not l.startswith("#")]
        assert len(packages) >= 25
        assert any(l.startswith("fastapi==") for l in packages)

    def test_deployment_files_exist(self, platform_dir):
        assert (platform_dir / "Dockerfile").is_file()
        assert (platform_dir / ".dockerignore").is_file()
        assert (platform_dir / "scripts" / "entrypoint.sh").is_file()
        assert (platform_dir / "scripts" / "serve.py").is_file()
        assert (platform_dir.parent / "docker-compose.yml").is_file()
        assert (platform_dir.parent / "DEPLOYMENT.md").is_file()

    def test_image_never_bakes_in_secrets(self, platform_dir):
        ignore = (platform_dir / ".dockerignore").read_text(encoding="utf-8")
        assert ".env" in ignore
        dockerfile = (platform_dir / "Dockerfile").read_text(encoding="utf-8")
        for secretish in ("PASSWORD=", "SECRET=", "postgresql://", "supabase"):
            assert secretish not in dockerfile

    def test_entrypoint_migrates_before_serving(self, platform_dir):
        text = (platform_dir / "scripts" / "entrypoint.sh").read_text(encoding="utf-8")
        assert text.index("alembic upgrade head") < text.index("exec uvicorn")

    def test_entrypoint_refuses_to_guess_a_database(self, platform_dir):
        text = (platform_dir / "scripts" / "entrypoint.sh").read_text(encoding="utf-8")
        assert "DATABASE_URL" in text and "exit 1" in text

    def test_container_defaults_are_the_safe_ones(self, platform_dir):
        text = (platform_dir / "Dockerfile").read_text(encoding="utf-8")
        assert "AUTO_CREATE_SCHEMA=false" in text
        assert "DEMO_SEED_ENABLED=false" in text
        assert "USER appuser" in text
